"""Выгрузка результатов: Excel и JSON."""
from __future__ import annotations

import json
import os

from .model import Drawing, SEVERITY_RU, ERROR, WARNING, INFO
from .tolerances import fmt

FILL = {ERROR: "F8D7DA", WARNING: "FFF3CD", INFO: "E7F1FB"}


def drawing_to_dict(d: Drawing) -> dict:
    tb = d.title
    return {
        "file": d.filename,
        "sheets": d.facts.get("sheets_label"),
        "kind": d.kind,
        "ocr": d.ocr,
        "error": d.error,
        "pages": [{"number": p.number, "width": p.width, "height": p.height, "image": p.image,
                   "scanned": p.scanned} for p in d.pages],
        "title": {"designation": tb.designation, "name": tb.name, "material": tb.material, "mass": tb.mass_raw,
                  "scale": tb.scale, "format": tb.format, "sheets": tb.sheets, "organization": tb.organization,
                  "people": tb.people, "changes": tb.changes, "found": tb.found},
        "dims": [{"raw": x.raw, "kind": x.kind, "nominal": x.nominal, "field": x.field,
                  "fit": f"{x.fit_hole}/{x.fit_shaft}" if x.fit_hole else None,
                  "upper": x.upper, "lower": x.lower, "reference": x.reference, "count": x.count,
                  "thread_tol": x.thread_tol, "page": x.page, "bbox": x.bbox} for x in d.dims],
        "tt": d.tt,
        "texts": d.facts.get("texts") or [ln.text for ln in d.lines][:2000],   # для правил предприятия
        "roughness": [{k: v for k, v in r.items() if k != "bbox"} for r in d.roughness],
        "limits": d.facts.get("limits", []),
        "facts": {k: d.facts.get(k) for k in ("overall", "sheet_thickness", "mass_bbox", "mass_ref",
                                              "is_assembly", "spec_item", "converted", "file_designation",
                                              "gear", "views", "spec_tree", "mass_view", "mass_parts", "spec_document", "spec_rows",
                                              "process", "classifier")},
        "findings": [f.to_dict() for f in d.findings],
        "suppressed": [f.to_dict() for f in d.suppressed],
        "welds": d.facts.get("welds", []),
        "dxf": d.facts.get("dxf"),
        "counts": d.counts,
    }


def save_json(drawings: list[Drawing], path: str) -> str:
    with open(path, "w", encoding="utf-8") as f:
        json.dump([drawing_to_dict(d) for d in drawings], f, ensure_ascii=False, indent=1, default=str)
    return path


def _as_dicts(drawings: list) -> list[dict]:
    return [d if isinstance(d, dict) else drawing_to_dict(d) for d in drawings]


def save_excel(drawings: list, path: str, completeness: dict | None = None) -> str:
    """drawings — объекты Drawing или словари drawing_to_dict (в т.ч. после применения исключений)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = _as_dicts(drawings)
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="34495E")
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, start=1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Сводка"
    header(ws, ["Файл", "Листы", "Обозначение", "Наименование", "Материал", "Масса", "Ошибки", "Предупр.",
                "Замечания", "Скрыто", "Итог"], [38, 9, 22, 26, 40, 9, 9, 9, 10, 8, 16])
    for d in data:
        c, t = d["counts"], d["title"]
        verdict = "Есть ошибки" if c[ERROR] else ("Проверить" if c[WARNING] else "Без замечаний")
        ws.append([d["file"], d.get("sheets") or "", t["designation"], t["name"], t["material"], t["mass"],
                   c[ERROR], c[WARNING], c[INFO], len(d.get("suppressed", [])), verdict])
        ws.cell(row=ws.max_row, column=11).fill = PatternFill(
            "solid", fgColor=FILL[ERROR] if c[ERROR] else (FILL[WARNING] if c[WARNING] else "D4EDDA"))
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Замечания")
    header(ws, ["Файл", "Обозначение", "Важность", "Раздел", "Код", "Замечание", "Стандарт", "Лист"],
           [34, 20, 15, 26, 7, 90, 22, 6])
    for d in data:
        for f in d["findings"]:
            ws.append([d["file"], d["title"]["designation"], SEVERITY_RU[f["severity"]], f["category"], f["code"],
                       f["message"], f.get("standard") or "", (f["page"] + 1) if f.get("page") is not None else ""])
            r = ws.max_row
            for col in range(1, 9):
                ws.cell(row=r, column=col).alignment = wrap
            ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=FILL[f["severity"]])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Размеры и допуски")
    header(ws, ["Файл", "Обозначение", "Размер", "Тип", "Номинал", "Поле/посадка", "Верхнее", "Нижнее",
                "Справочный", "Кол-во"], [34, 20, 26, 11, 10, 14, 10, 10, 11, 8])
    kinds = {"linear": "линейный", "diameter": "диаметр", "radius": "радиус", "thread": "резьба",
             "chamfer": "фаска", "angle": "угол", "thickness": "толщина", "sphere": "сфера", "square": "квадрат"}
    for d in data:
        for x in d["dims"]:
            ws.append([d["file"], d["title"]["designation"], x["raw"], kinds.get(x["kind"], x["kind"]), x["nominal"],
                       x["field"] or x["fit"] or "",
                       fmt(x["upper"], True) if x["upper"] is not None else "",
                       fmt(x["lower"], True) if x["lower"] is not None else "", x["reference"] or "",
                       x["count"] or ""])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Технические требования")
    header(ws, ["Файл", "Обозначение", "Пункт ТТ"], [34, 20, 110])
    for d in data:
        for t in d["tt"]:
            ws.append([d["file"], d["title"]["designation"], t])
            ws.cell(row=ws.max_row, column=3).alignment = wrap

    if completeness:
        ws = wb.create_sheet("Комплектность", 1)
        header(ws, ["Поз.", "Обозначение", "Артикул", "Наименование", "Раздел", "Нужен чертёж", "Статус", "Файл",
                    "Ошибки", "Предупр."], [8, 22, 13, 32, 18, 12, 18, 40, 8, 9])
        names = {"ok": "Без замечаний", "warnings": "Проверить", "errors": "Есть ошибки", "missing": "НЕТ ЧЕРТЕЖА",
                 "not_required": "не требуется"}
        fills = {"ok": "D4EDDA", "warnings": FILL[WARNING], "errors": FILL[ERROR], "missing": "F5C2C7"}
        for r in completeness["rows"]:
            ws.append([r["pos"], r["designation"], r["sku"], r["name"], r["section"], "да" if r["need"] else "нет",
                       names[r["status"]], r["file"], r["errors"] or "", r["warnings"] or ""])
            if r["status"] in fills:
                ws.cell(row=ws.max_row, column=7).fill = PatternFill("solid", fgColor=fills[r["status"]])
        for x in completeness["extra"]:
            ws.append(["", x["designation"], "", "", "", "", "нет в спецификации", x["file"], "", ""])
        ws.auto_filter.ref = ws.dimensions

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
