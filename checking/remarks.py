"""Перечень замечаний для конструктора (Word), по образцу перечня замечаний нормоконтролёра ГОСТ 2.111."""
from __future__ import annotations

import time

from .model import SEVERITY_RU


def build_docx(drawings: list[dict], path: str, include_info: bool = False, title: str | None = None) -> str:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(15))
    sec.top_margin = sec.bottom_margin = Mm(12)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    h = doc.add_paragraph()
    r = h.add_run(title or "Перечень замечаний по проверке чертежей")
    r.bold = True
    r.font.size = Pt(14)
    docs = [d for d in drawings if d.get("findings")]
    p = doc.add_paragraph(f"Дата: {time.strftime('%d.%m.%Y')}    Документов: {len(drawings)}    "
                          f"с замечаниями: {len(docs)}")
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x5E, 0x6E)

    headers = ["№", "Обозначение документа", "Лист", "Содержание замечания", "Ссылка на НД", "Важность",
               "Отметка об устранении"]
    widths = [Mm(10), Mm(38), Mm(12), Mm(118), Mm(32), Mm(22), Mm(35)]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (c, txt) in enumerate(zip(table.rows[0].cells, headers)):
        c.text = ""
        run = c.paragraphs[0].add_run(txt)
        run.bold = True
        run.font.size = Pt(9)
        c.width = widths[i]
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8ECF2")
        c._tc.get_or_add_tcPr().append(shd)
    # повтор шапки на каждой странице
    trPr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)

    n = 0
    colors = {"error": "FDECEC", "warning": "FFF4E0", "info": "E9F1FB"}
    for d in drawings:
        items = [f for f in d.get("findings", []) if include_info or f["severity"] != "info"]
        if not items:
            continue
        des = (d.get("title") or {}).get("designation") or d.get("file", "")
        name = (d.get("title") or {}).get("name") or ""
        for f in items:
            n += 1
            row = table.add_row().cells
            vals = [str(n), f"{des}\n{name}".strip(), str(f["page"] + 1) if f.get("page") is not None else "",
                    f["message"], f.get("standard") or "", SEVERITY_RU.get(f["severity"], f["severity"]), ""]
            for i, (c, v) in enumerate(zip(row, vals)):
                c.text = v
                c.width = widths[i]
                for par in c.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), colors[f["severity"]])
            row[5]._tc.get_or_add_tcPr().append(shd)
    if n == 0:
        doc.add_paragraph("Замечаний нет.")

    doc.add_paragraph()
    for line in ("Проверил (нормоконтроль): ____________________ / ____________________   «___» __________ 20__ г.",
                 "С замечаниями ознакомлен (разработчик): ____________________ / ____________________",
                 "Замечания устранены: ____________________ / ____________________   «___» __________ 20__ г."):
        doc.add_paragraph(line)
    doc.save(path)
    return path
