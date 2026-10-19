"""Маршрутные карты и технологические процессы (ЕСТД): приём файла, распознавание вида документа, обозначения детали и
перечня операций. Правил проверки пока нет — документ показывается в отчёте и связывается с чертежом по обозначению.

Читаются PDF (текст или скан через OCR), Excel (xlsx/xls) и Word (docx). Распознавание ведётся по тексту, поэтому
для нестандартных форм предприятия список операций может оказаться неполным — это уточняется по образцам."""
from __future__ import annotations

import os
import re
from collections import Counter

from .model import Drawing, INFO, CAT_FILE
from .textnorm import nfc, norm_designation

ABBR = {"MK": "МК", "OK": "ОК"}
PROCESS_EXT = {".pdf", ".xlsx", ".xlsm", ".xls", ".docx"}

# вид документа: (код, название, регулярное выражение по тексту первых листов)
KINDS = [
    ("MK", "Маршрутная карта (МК)", r"маршрутн\w*\s+карт|маршрутно-технологическ\w*\s+карт"),
    ("OK", "Операционная карта (ОК)", r"операционн\w*\s+карт"),
    ("KE", "Карта эскизов (КЭ)", r"карт\w*\s+эскиз"),
    ("KTK", "Карта технического контроля (КТК)", r"карт\w*\s+технического\s+контрол"),
    ("TI", "Технологическая инструкция (ТИ)", r"технологическ\w*\s+инструкци"),
    ("VT", "Ведомость технологических документов", r"ведомост\w*\s+(?:технологическ|оснастки|материалов|деталей)"),
    ("TL", "Титульный лист комплекта", r"комплект\w*\s+документов\s+на\s+технологическ|титульн\w*\s+лист"),
    ("TP", "Технологический процесс", r"технологическ\w*\s+процесс"),
]
OPERATIONS = ("заготовит", "токарн", "фрезерн", "сверлильн", "расточн", "шлифовальн", "протяжн", "зубофрезерн", "зубошлифов",
              "зубодолбежн", "резьбонарезн", "слесарн", "сварочн", "термическ", "термообработ", "гальваническ", "покрытие",
              "контрольн", "маркировоч", "упаковоч", "лазерн", "плазменн", "гибочн", "штамповоч", "отрезн", "малярн",
              "окрасочн", "полировальн", "пескоструйн", "дробеструйн", "моечн", "сборочн", "балансировоч", "электроэрозион",
              "хонинговальн", "координатно-расточн", "притирочн", "правильн", "обработка резанием", "ковочн", "литейн",
              "зачистн", "вальцовочн", "кузнечн", "транспортн", "консервац", "травлен", "анодирован", "оксидирован")
DESIGNATION = re.compile(r"[А-ЯA-Z]{2,5}[.\s]?\d{6}\.\d{3}(?:-\d{2,3})?|\b\d{3}-\d{6}\b|[А-ЯA-Z]{1,4}-\d{4}\.\d{3}")
OP_NO = re.compile(r"(?<![\d.,])(\d{3})(?![\d.,])")
OP_NAME = re.compile(r"(?i)\b\w*(?:" + "|".join(OPERATIONS) + r")\w*(?:\s+\w+)?")
EQUIP = re.compile(r"(?i)\b(?:станок|центр|пресс|печь|установка|машина|автомат|верстак|стенд|гибочный\s+пресс)[^|;]{0,40}")


def _pdf_lines(path: str, img_dir: str | None, use_ocr: bool) -> tuple[list[str], Drawing]:
    from .pdfreader import read_pdf
    drw = Drawing(path, os.path.basename(path), kind="pdf")
    read_pdf(path, drw, img_dir, use_ocr=use_ocr)
    return [ln.text.strip() for ln in drw.lines if ln.text.strip()], drw


def _xlsx_lines(path: str) -> list[str]:
    out = []
    if path.lower().endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(path, logfile=open(os.devnull, "w"))
        for sh in wb.sheets():
            for r in range(sh.nrows):
                cells = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
                if any(cells):
                    out.append(" | ".join(c for c in cells if c))
        return out
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [("" if v is None else str(v).strip()) for v in row]
            if any(cells):
                out.append(" | ".join(c for c in cells if c))
    wb.close()
    return out


def _docx_lines(path: str) -> list[str]:
    import docx
    d = docx.Document(path)
    out = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells, seen = [], set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                cells.append(c.text.strip().replace("\n", " "))
            if any(cells):
                out.append(" | ".join(c for c in cells if c))
    return out


def _labeled(ln: str, label: str) -> str:
    """Значение после подписи в той же ячейке («Материал: Сталь 45») или в следующей («Материал | Сталь 45»)."""
    cells = [c.strip() for c in ln.split("|")]
    for i, c in enumerate(cells):
        m = re.match(label + r"\s*[:]?\s*(.*)$", c, re.I)
        if m:
            rest = m.group(1).strip()
            for v in ([rest] if rest else []) + cells[i + 1:]:
                if v and not DESIGNATION.fullmatch(v) and not re.fullmatch(r"[\d\s.,-]+", v)                         and not re.match(r"(?i)(?:основной|наименование|код)", v):
                    return v[:70].strip(" :;")
    return ""


def recognize(lines: list[str]) -> dict:
    """Вид документа, обозначение детали, наименование, материал и операции по строкам текста."""
    head, full = chr(10).join(lines[:60]), chr(10).join(lines)
    kind = None
    for text in (head, full):                          # сначала шапка документа, затем весь текст
        for k, n, rx in KINDS:
            if re.search(rx, text, re.I):
                kind = (k, n)
                break
        if kind:
            break
    if kind is None:                                   # только сокращение в шапке: «МК», «ОК»
        kind = next(((k, n) for k, n, _ in KINDS if k in ABBR and re.search(r"(?<![А-Яа-я])" + ABBR[k] + r"(?![А-Яа-я])", head)), None)
    des = Counter()
    for ln in lines:
        for m in DESIGNATION.finditer(ln):
            des[re.sub(r"\s+", "", m.group(0))] += 1
    designation = des.most_common(1)[0][0] if des else ""
    name = material = ""
    for ln in lines:
        name = name or _labeled(ln, r"(?:наименование(?:\s+(?:детали|изделия|сб\.?\s*единицы|сборочной\s+единицы))?|деталь)")
        material = material or _labeled(ln, r"материал\w*")
        if name and material:
            break
    ops, seen = [], set()
    for ln in lines:
        m = OP_NAME.search(ln)
        if not m:
            continue
        nums = [x for x in OP_NO.findall(ln[: m.start() + 1]) if int(x) % 5 == 0 and 0 < int(x) < 1000]
        no = nums[-1] if nums else ""
        if no and no in seen:
            continue
        seen.add(no)
        eq = EQUIP.search(ln)
        ops.append({"no": no, "name": m.group(0).strip().capitalize(), "equipment": eq.group(0).strip() if eq else "",
                    "line": ln[:160]})
    return {"type": kind[0] if kind else "", "type_name": kind[1] if kind else "Документ технологии (вид не определён)",
            "designation": designation, "name": name, "material": material, "operations": ops[:200]}


def check_process_file(path: str, cfg: dict | None = None, img_dir: str | None = None,
                       display_name: str | None = None) -> Drawing:
    """Читает маршрутную карту / техпроцесс и возвращает документ отчёта (kind = «process»)."""
    cfg = cfg or {}
    name = nfc(display_name or os.path.basename(path))
    ext = os.path.splitext(path)[1].lower()
    drw = Drawing(path, name, kind="process")
    try:
        if ext == ".pdf":
            if img_dir:
                os.makedirs(img_dir, exist_ok=True)
            lines, src = _pdf_lines(path, img_dir, cfg.get("ocr", True))
            drw.pages, drw.ocr, drw.lines = src.pages, src.ocr, src.lines
        elif ext in (".xlsx", ".xlsm", ".xls"):
            lines = _xlsx_lines(path)
        elif ext == ".docx":
            lines = _docx_lines(path)
        else:
            lines = []
    except Exception as e:
        drw.error = f"{type(e).__name__}: {e}"
        drw.add("error", "F00", CAT_FILE, f"Файл не удалось прочитать: {drw.error}")
        return drw
    info = recognize(lines)
    drw.title.found = True
    drw.title.designation = info["designation"]
    drw.title.name = info["name"]
    drw.title.material = info["material"]
    drw.facts["process"] = info
    drw.facts["file_designation"] = info["designation"]
    drw.facts["texts"] = lines[:5000]
    if not lines:
        drw.add(INFO, "F08", CAT_FILE, "В документе не найдено текста: для сканов включите распознавание (OCR)")
    elif not info["type"] and not info["operations"]:
        drw.add(INFO, "F08", CAT_FILE, "Документ не распознан как маршрутная карта или технологический процесс — "
                "вид документа и операции не определены")
    return drw


def link_to_drawings(docs: list) -> None:
    """Для документов технологии находит чертёж детали с тем же обозначением среди документов проверки."""
    index = {}
    for d in docs:
        if d.kind in ("process", "spec") or d.error:
            continue
        for x in (d.title.designation, d.facts.get("file_designation")):
            if x:
                index.setdefault(norm_designation(re.sub(r"\s*(СБ|МЧ|ВО|ГЧ)\s*$", "", x)), d)
    for d in docs:
        p = d.facts.get("process")
        if d.kind != "process" or not p:
            continue
        hit = index.get(norm_designation(p.get("designation") or "")) if p.get("designation") else None
        p["drawing"] = hit.filename if hit is not None else None
