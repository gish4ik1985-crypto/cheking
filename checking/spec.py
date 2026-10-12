"""Загрузка спецификации / выгрузки ERP (xlsx, xlsm, xls) и поиск позиции по обозначению.

Поддерживаются:
* отчёт КОМПАС «Report» (N, Наименование, Раздел спецификации, Обозначение, Количество, Материал, Масса);
* шаблон СП ERP ЯСТ (Артикул, Обозначение децимальное, Наименование, …, Кол-во на ед. шт. (вес…),
  Полное наименование материалов, Марка, Тип/Базовый размер/Размеры заготовки);
* табличный экспорт ERP (Арт./Дец. номер/Наименование/Масса ед./Наим. материала/…).
"""
from __future__ import annotations

import os
import re

from .textnorm import norm_designation, to_float

FIELDS = ("designation", "name", "material", "grade", "mass", "sku", "base_size", "blank_size",
          "blank_type", "section", "pos", "note")


def fmt_pos(v) -> str:
    """Номер позиции: 1.0 → «1», «1.1.2» как есть."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _hdr(v) -> str:
    return " ".join(str(v or "").lower().replace("\n", " ").split())


def _map_columns(header: list) -> dict:
    cols: dict[str, int] = {}
    for i, h in enumerate(map(_hdr, header)):
        if not h:
            continue
        parent = any(x in h for x in ("сбор", "родит"))
        if h.startswith("итого"):
            continue
        if h.startswith("кол"):                      # «Кол-во, шт. на одну сборочную единицу», «Количество»
            if "вес" in h or "объем" in h:
                cols.setdefault("mass", i)            # шаблон ERP: вес/объём/площадь на единицу
                cols.setdefault("amount", i)
            else:
                cols.setdefault("qty", i)
            continue
        if h.startswith("ед.") or h.startswith("ед ") or "ед. изм" in h:
            cols.setdefault("unit", i)
            continue
        if parent:
            if "артикул" in h:
                cols.setdefault("parent_sku", i)
            elif "обозначение" in h:
                cols.setdefault("parent_des", i)
            elif "наименование" in h:
                cols.setdefault("parent_name", i)
            continue
        if ("обозначение" in h or h.startswith("дец")) and not parent:
            cols.setdefault("designation", i)
        elif "наименование" in h and not parent and "материал" not in h and "наим." not in h:
            cols.setdefault("name", i)
        elif h.startswith("наим") and "материал" in h:
            cols.setdefault("material", i)
        elif "материал" in h and "артикул" not in h and "арт." not in h and "марка" not in h:
            cols.setdefault("material", i)
        elif "марка" in h:
            cols.setdefault("grade", i)
        elif "масса" in h or ("вес" in h and "итого" not in h):
            cols.setdefault("mass", i)
        elif ("артикул" in h or h.startswith("арт")) and not parent and "материал" not in h:
            cols.setdefault("sku", i)
        elif "базовый размер" in h:
            cols.setdefault("base_size", i)
        elif "размеры заготовки" in h or "размер заготовки" in h:
            cols.setdefault("blank_size", i)
        elif "тип заготовки" in h:
            cols.setdefault("blank_type", i)
        elif "раздел" in h:
            cols.setdefault("section", i)
        elif h in ("n", "№", "поз.", "поз") or h.startswith("№ поз"):
            cols.setdefault("pos", i)
        elif h.startswith("примеч"):
            cols.setdefault("note", i)
    return cols


def _rows(path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path, logfile=open(os.devnull, "w"))
        sh = wb.sheet_by_index(0)
        for r in range(sh.nrows):
            yield [sh.cell_value(r, c) for c in range(sh.ncols)]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            yield list(row)
        wb.close()


class Spec:
    def __init__(self, path: str):
        self.path = path
        self.items: dict[str, dict] = {}
        self.by_sku: dict[str, dict] = {}
        self.rows: list[dict] = []
        self.columns: dict[str, int] = {}
        self._load()

    def _load(self):
        header_found = False
        for row in _rows(self.path):
            if not header_found:
                cols = _map_columns(row)
                if "designation" in cols and ("name" in cols or "mass" in cols):
                    self.columns = cols
                    header_found = True
                continue
            item = {f: (row[i] if i < len(row) else None) for f, i in self.columns.items()}
            des = str(item.get("designation") or "").strip()
            sku = str(item.get("sku") or "").strip()
            if not des and not sku and not (str(item.get("name") or "").strip() and item.get("pos") not in (None, "")):
                continue                              # стандартные/прочие изделия без обозначения — по позиции
            # выгрузки, где в графе «Обозначение» стоит артикул (241-000779), — артикул переносим
            if not sku and re.fullmatch(r"\d{3}-\d{6}", des):
                item["sku"] = sku = des
                item["designation_is_sku"] = True
            item["designation"] = des
            item["mass"] = item["mass"] if isinstance(item.get("mass"), (int, float)) else to_float(
                str(item.get("mass") or "").replace(" ", "") or None)
            q = item.get("qty")
            item["qty"] = q if isinstance(q, (int, float)) else to_float(str(q or "").replace(" ", "") or None)
            for k in ("name", "material", "grade", "sku", "base_size", "blank_size", "blank_type", "section",
                      "pos", "note", "unit", "parent_sku", "parent_des", "parent_name"):
                v = item.get(k)
                item[k] = "" if v is None else (fmt_pos(v) if k == "pos" else str(v).strip())
            item["row"] = len(self.rows) + 1
            self.rows.append(item)
            if des:
                self.items.setdefault(norm_designation(des), item)
            if sku:
                self.by_sku.setdefault(sku, item)
        if not header_found:
            raise ValueError("Не найдена строка заголовков со столбцом «Обозначение»")

    def find(self, designation: str, sku: str | None = None) -> dict | None:
        hit = self.items.get(norm_designation(designation)) if designation else None
        if hit is None and sku:
            hit = self.by_sku.get(sku)
        return hit

    def __len__(self):
        return max(len(self.items), len(self.by_sku))


class ItemSpec:
    """«Спецификация» из одной позиции — данные карточки изделия, переданные из ERP в API-запросе."""

    def __init__(self, item: dict, source: str = "ERP"):
        self.path = source
        self.item = {f: item.get(f) for f in FIELDS}
        self.item["designation"] = str(item.get("designation") or "")
        self.item["sku"] = str(item.get("sku") or "")
        m = item.get("mass")
        self.item["mass"] = to_float(str(m)) if m not in (None, "") else None
        for k in ("name", "material", "grade", "base_size", "blank_size", "blank_type", "section"):
            self.item[k] = str(item.get(k) or "")
        self.item["dxf_length"] = to_float(str(item.get("dxf_length") or "")) if item.get("dxf_length") else None
        self.item["dxf_width"] = to_float(str(item.get("dxf_width") or "")) if item.get("dxf_width") else None

    def find(self, designation: str, sku: str | None = None) -> dict:
        return self.item           # запрос всегда про конкретное изделие

    def __len__(self):
        return 1
