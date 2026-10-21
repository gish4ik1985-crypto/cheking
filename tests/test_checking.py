"""Тесты: таблицы допусков, разбор размеров, правила на тестовых и реальных чертежах.

Запуск:  .venv\\Scripts\\python -m pytest -q
"""
from __future__ import annotations

import os
import subprocess
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from checking import tolerances as T  # noqa: E402
from checking.dimparse import parse_dimension  # noqa: E402
from checking.engine import check_file  # noqa: E402
from checking.material import parse_material  # noqa: E402
from checking.spec import Spec  # noqa: E402
from checking.standards import check_metric_thread  # noqa: E402

SAMPLES = os.path.join(ROOT, "samples")


# --- ГОСТ 25346 / ISO 286: контрольные значения из таблиц ISO 286-2 ----------
@pytest.mark.parametrize("size,code,up,lo", [
    (20, "H7", 0.021, 0), (20, "g6", -0.007, -0.020), (20, "f7", -0.020, -0.041), (50, "k6", 0.018, 0.002),
    (40, "p6", 0.042, 0.026), (25, "r6", 0.041, 0.028), (100, "s6", 0.093, 0.071), (30, "K7", 0.006, -0.015),
    (30, "N7", -0.007, -0.028), (30, "P7", -0.014, -0.035), (30, "M7", 0, -0.021), (60, "e8", -0.060, -0.106),
    (80, "d9", -0.100, -0.174), (10, "js6", 0.0045, -0.0045), (50, "F8", 0.064, 0.025), (40, "G7", 0.034, 0.009),
    (100, "S7", -0.058, -0.093), (160, "d11", -0.145, -0.395), (200, "t6", 0.195, 0.166), (20, "K6", 0.002, -0.011),
    (20, "N9", 0, -0.052), (20, "M8", 0.004, -0.029), (20, "J7", 0.012, -0.009), (2, "N7", -0.004, -0.014),
    (35, "u6", 0.076, 0.060), (15, "x7", 0.063, 0.045), (120, "H7", 0.035, 0), (1000, "h9", 0, -0.230),
])
def test_limits(size, code, up, lo):
    got = T.limits(size, code)
    assert got == pytest.approx((up, lo), abs=1e-6)


def test_fit_kinds():
    assert T.fit_info(20, "H7", "g6").kind == "с зазором"
    assert T.fit_info(25, "H7", "k6").kind == "переходная"
    assert T.fit_info(40, "H7", "s6").kind == "с натягом"


@pytest.mark.parametrize("code", ["Js7", "H19", "Q7", "h0"])
def test_bad_fields(code):
    with pytest.raises(T.ToleranceError):
        T.limits(30, code)


def test_field_not_defined_for_size():
    with pytest.raises(T.ToleranceError):
        T.limits(12, "t6")          # t начинается с 24 мм


# --- резьба --------------------------------------------------------------------
def test_threads():
    assert check_metric_thread(18, None)[0]
    assert check_metric_thread(20, 1.5)[0]
    assert not check_metric_thread(13, None)[0]
    assert not check_metric_thread(20, 1.25)[0]
    assert not check_metric_thread(40, None)[0]      # М40 — только мелкий шаг


# --- разбор текста размеров ------------------------------------------------------
@pytest.mark.parametrize("text,kind,nom,extra", [
    ("⌀20H7", "diameter", 20, {"field": "H7"}),
    ("⌀40H7/g6", "diameter", 40, {"fit_hole": "H7", "fit_shaft": "g6"}),
    ("110±1", "linear", 110, {"upper": 1, "lower": -1, "symmetric": True}),
    ("⌀20H7(+0,021)", "diameter", 20, {"field": "H7", "upper": 0.021, "lower": 0}),
    ("M18×1,5-6g", "thread", 18, {"pitch": 1.5, "thread_tol": "6g"}),
    ("2×45°", "chamfer", 2, {"angle": 45}),
    ("4 отв. ⌀10", "diameter", 10, {"count": 4}),
    ("16*", "linear", 16, {"reference": "*"}),
    ("R3**", "radius", 3, {"reference": "**"}),
    ("s3*", "thickness", 3, {"reference": "*"}),
    ("⌀45 Н7", "diameter", 45, {"field": "H7", "cyrillic_in_field": True}),
])
def test_parse_dimension(text, kind, nom, extra):
    d = parse_dimension(text)
    assert d and d["kind"] == kind and d["nominal"] == pytest.approx(nom)
    for k, v in extra.items():
        assert d.get(k) == (pytest.approx(v) if isinstance(v, float) else v), k


@pytest.mark.parametrize("text", ["1:2", "19.08.24", "3-2025", "301321.001", "01.2025", "Ст3сп", "А-А", "Ra 6,3"])
def test_not_dimension(text):
    assert parse_dimension(text) is None


def test_material():
    m = parse_material("Лист Б-ПН-О-16х1500х6000 ГОСТ 19903-2015 / Ст3сп-5-ГС ГОСТ 14637-2024")
    assert (m["sortament"], m["size"], m["grade"]) == ("Лист", 16, "Ст3сп")
    m = parse_material("Профиль 200х120х6 ГОСТ 30245-2003 / 345-14-09Г2С ГОСТ 19281-2014")
    assert m["grade"] == "09Г2С"
    m = parse_material("Круг В1-IV-НД-100 ГОСТ 2590-2006/20-2ГП-М1-ТВ1 ГОСТ 1050-2013")
    assert (m["size"], m["grade"]) == (100, "20")


# --- тестовые чертежи с заложенными ошибками ------------------------------------
@pytest.fixture(scope="module")
def generated():
    subprocess.run([sys.executable, os.path.join(ROOT, "tests", "make_samples.py")], check=True,
                   capture_output=True)
    return os.path.join(SAMPLES, "test")


def _mfg_cfg():
    """Настройки с включённым блоком «Технологичность» (по умолчанию он выключен)."""
    from checking.config import load_config
    return {**load_config(), "modules": {"mfg": True}}


def _codes(path, spec=None, cfg=None):
    ds = check_file(path, spec=spec, cfg=cfg)
    assert len(ds) == 1
    return ds[0], {f.code for f in ds[0].findings}


def test_errors_drawing(generated):
    d, codes = _codes(os.path.join(generated, "ДТМЛ.741111.900 - Планка тестовая.pdf"), cfg=_mfg_cfg())
    for c in ("D01", "D03", "D04", "D05", "D06", "D08", "D11", "D13", "D14", "D15", "D29", "T01", "T02", "T03",
              "D10", "B03"):
        assert c in codes, c
    msgs = " ".join(f.message for f in d.findings)
    assert "+0,021 / 0" in msgs                          # правильные отклонения для ⌀20H7
    assert "H7/g6" in msgs                               # подсказка для перевёрнутой посадки


def test_clean_drawing(generated):
    d, codes = _codes(os.path.join(generated, "ДТМЛ.715123.901 - Вал.pdf"))
    errors = [f for f in d.findings if f.severity == "error"]
    assert not errors, [f.message for f in errors]
    assert "D10" not in codes                            # общие допуски указаны в ТТ


# --- реальные чертежи КОМПАС + спецификация ---------------------------------------
REAL = os.path.join(SAMPLES, "drawings")


@pytest.mark.skipif(not os.path.isdir(REAL), reason="нет образцов")
def test_real_mass_error_found():
    spec = Spec(os.path.join(SAMPLES, "СП Опора пассивная.xls"))
    d, codes = _codes(os.path.join(REAL, "ДТМЛ.741224.002 - Фиксатор изм.2.pdf"), spec)
    assert "B02" in codes                                # 0,05 кг в штампе против 0,118 кг в СП


@pytest.mark.skipif(not os.path.isdir(REAL), reason="нет образцов")
def test_real_clean_drawings_have_no_errors():
    spec = Spec(os.path.join(SAMPLES, "СП Опора пассивная.xls"))
    for name in ("ДТМЛ.741214.001 - Пластина изм.1.pdf", "ДТМЛ.741228.001 - Основание.pdf",
                 "ДТМЛ.711111.001 - Кольцо.pdf", "ДТМЛ.715111.001 - Шпилька.pdf"):
        d, _ = _codes(os.path.join(REAL, name), spec)
        errors = [f.message for f in d.findings if f.severity == "error"]
        assert not errors, (name, errors)
        assert d.title.designation.startswith("ДТМЛ.")


@pytest.mark.skipif(not os.path.isdir(REAL), reason="нет образцов")
def test_real_dimension_between_extension_lines():
    # «25» и «20*» между выносными линиями под контуром — размеры, а не ячейки таблицы;
    # знак × в «5×45°» без ToUnicode в шрифте GOSTTypeA
    d, codes = _codes(os.path.join(REAL, "ДТМЛ.741124.068 - Пластина.pdf"))
    raws = {x.raw: x for x in d.dims}
    assert "25" in raws and raws["20*"].reference == "*" and raws["5×45°"].kind == "chamfer"
    assert "D09" not in codes


# --- реальные чертежи Solid Edge (редуктор), найденные ошибки подтверждены по изображению ----
RED = os.path.join(SAMPLES, "reductor")


def _find(d, code, text):
    return [f for f in d.findings if f.code == code and text in f.message]


@pytest.mark.skipif(not os.path.isdir(RED), reason="нет образцов редуктора")
def test_reductor_real_findings():
    d, _ = _codes(os.path.join(RED, "201-000571 Палец сателлита диам 30 мм версия 2.pdf"))
    assert d.title.designation == "201-000571"
    assert _find(d, "D04", "«30 -0,2 / -0,1»") and _find(d, "D04", "⌀28,5")       # отклонения перепутаны
    assert _find(d, "D14", "h6")                                                   # h6 при общей Ra 6,3
    d, _ = _codes(os.path.join(RED, "280-000091 Водило 2 ступени версия 4.pdf"))
    assert _find(d, "D04", "⌀168")                                                 # −0,05 / −0,05
    assert _find(d, "D03", "74H8")                                                 # +0,05 вместо +0,046
    d, _ = _codes(os.path.join(RED, "280-000092 Водило 3 ступени версия 4.pdf"))
    assert _find(d, "D04", "⌀100 +1 / +1,5")


@pytest.mark.skipif(not os.path.isdir(RED), reason="нет образцов редуктора")
def test_reductor_layout():
    # страница с /Rotate 90 и основная надпись «Си Джи Эм» с графами Рев./Вер.
    d, codes = _codes(os.path.join(RED, "204-000255 Сателлит водила  2 ступени.pdf"))
    assert d.title.designation == "204-000255" and d.title.mass == 2.31
    # ТТ под таблицей спецификации на сборочном чертеже
    d, codes = _codes(os.path.join(RED, "120-001422 Сателлит 3 ступени в сборе с иголч подш.pdf"))
    assert "D10" not in codes
    # числа из таблицы параметров зубчатого венца — не размеры
    d, _ = _codes(os.path.join(RED, "280-000059 Зубчатый венец m3 z75.pdf"))
    assert not [x for x in d.dims if x.raw in ("5.176", "8.69")]


# --- DXF -----------------------------------------------------------------------
def test_dxf_compare(generated):
    from checking.dxf import DxfIndex
    idx = DxfIndex([os.path.join(generated, "ДТМЛ.741111.900 - Планка тестовая_12мм.dxf"),
                    os.path.join(generated, "ДТМЛ.741214.001 Пластина_16мм.dxf")])
    ds = check_file(os.path.join(generated, "ДТМЛ.741111.900 - Планка тестовая.pdf"), dxf_index=idx)
    msgs = [f.message for f in ds[0].findings if f.code.startswith("X")]
    assert any("4 отв., в DXF 3" in m for m in msgs)
    assert any("⌀12,5 — 2 шт." in m for m in msgs)
    ds = check_file(os.path.join(REAL, "ДТМЛ.741214.001 - Пластина изм.1.pdf"), dxf_index=idx)
    assert [f.code for f in ds[0].findings if f.code.startswith("X")] == ["X00"]


# --- HTTP API для ERP -------------------------------------------------------------
@pytest.mark.skipif(not os.path.isdir(REAL), reason="нет образцов")
def test_api_check_with_erp_card():
    import app as webapp
    client = webapp.app.test_client()
    assert client.get("/api/v1/health").json["status"] == "ok"
    with open(os.path.join(REAL, "ДТМЛ.741224.002 - Фиксатор изм.2.pdf"), "rb") as f:
        r = client.post("/api/v1/check", data={
            "file": (f, "ДТМЛ.741224.002 - Фиксатор изм.2.pdf"), "designation": "ДТМЛ.741224.002",
            "mass": "0.1178", "material": "Лист 5 ГОСТ 19903-2015 / Ст3сп ГОСТ 14637-89", "ocr": "0"},
            content_type="multipart/form-data")
    assert r.status_code == 200 and r.json["success"]
    d = r.json["drawings"][0]
    assert any(f["code"] == "B02" and "карточке ERP" in f["message"] for f in d["findings"])


# --- комплектность, сварка, исключения, перечень замечаний ---------------------
@pytest.mark.skipif(not os.path.isdir(REAL), reason="нет образцов")
def test_completeness():
    from checking.completeness import completeness
    from checking.engine import collect_files
    spec = Spec(os.path.join(SAMPLES, "СП Опора пассивная.xls"))
    ds = [d for f in collect_files([REAL]) for d in check_file(f, spec=spec)]
    c = completeness(spec, ds)
    missing = {r["designation"] for r in c["rows"] if r["status"] == "missing"}
    assert "ДТМЛ.715123.001" in missing and "ДТМЛ.741214.001" not in missing   # Вал — нет, Пластина — есть
    assert any(r["designation"] == "ДТМЛ.741224.002" and r["status"] == "errors" for r in c["rows"])


def test_weld_parse():
    from checking.welds import _parse
    w = _parse("ГОСТ 14771-76-Т3-ИП-⊿5")
    assert (w["std"], w["type"], w["num"], w["leg"]) == ("14771", "Т", 3, 5.0)
    assert _parse("У5-⊿")["leg"] is None
    assert _parse("С17") is None


def test_exceptions_roundtrip(tmp_path):
    from checking import exceptions as ex
    p = str(tmp_path / "ex.json")
    d = {"file": "a.pdf", "title": {"designation": "ДТМЛ.1"},
         "findings": [{"code": "D10", "message": "m1", "severity": "warning"},
                      {"code": "B01", "message": "m2", "severity": "warning"}]}
    r = ex.add_rule({"scope": "drawing", "code": "D10", "message": "m1", "designation": "ДТМЛ.1"}, p)
    out = ex.apply_to_dict(d, ex.load_rules(p))
    assert [f["code"] for f in out["findings"]] == ["B01"] and out["counts"]["warning"] == 1
    ex.add_rule({"scope": "global", "code": "B01"}, p)
    assert ex.apply_to_dict(d, ex.load_rules(p))["findings"] == []
    assert ex.delete_rule(r["id"], p)


def test_remarks_docx(tmp_path, generated):
    from docx import Document
    from checking.remarks import build_docx
    from checking.report import drawing_to_dict
    ds = check_file(os.path.join(generated, "ДТМЛ.741111.900 - Планка тестовая.pdf"))
    path = build_docx([drawing_to_dict(d) for d in ds], str(tmp_path / "r.docx"))
    rows = Document(path).tables[0].rows
    assert len(rows) > 10 and "ДТМЛ.741111.900" in rows[1].cells[1].text


# --- проверки по остальным ГОСТам на реальных чертежах (находки подтверждены по изображениям) -----
def _red(name):
    return check_file(os.path.join(RED, name))[0]


@pytest.mark.skipif(not os.path.isdir(RED), reason="нет образцов редуктора")
def test_gost_findings_reductor():
    d = _red("201-000571 Палец сателлита диам 30 мм версия 2.pdf")
    assert _find(d, "K01", "С4")                                     # форма C — только d 8, 12, 20…
    assert _find(d, "P04", "IT14/2")                                 # без «±»
    d = _red("280-000060 Зубчатый венец редуктора Прево 2 ступени.pdf")
    assert _find(d, "G01", "«Г»")                                    # база Г = разрез Г-Г, база — Д
    assert _find(d, "G04", "«Д»")
    assert _find(d, "N01", "13755-81") and _find(d, "Z04", "8C")
    assert not [f for f in d.findings if f.code == "H02"]            # «слоя 0.4...0.5 HRC 52...58» — не твёрдость
    d = _red("204-000255 Сателлит водила  2 ступени.pdf")
    assert _find(d, "M01", "Круг") and _find(d, "Z05", "ГОСТ 1643")
    assert not [f for f in d.findings if f.code == "N02"]            # номер стандарта перенесён на строку ниже
    d = _red("280-000091 Водило 2 ступени версия 4.pdf")
    assert _find(d, "K04", "6033-88")
    d = _red("280-000092 Водило 3 ступени версия 4.pdf")
    assert _find(d, "K04", "6333-80") and not [f for f in d.findings if f.code == "Z02"]   # таблица шлицев
    d = _red("280-000059 Зубчатый венец m3 z75.pdf")
    assert d.facts["gear"][0]["m"] == "3" and not [f for f in d.findings if f.code in ("Z03", "G01")]


@pytest.mark.skipif(not os.path.isdir(REAL), reason="нет образцов")
def test_gost_findings_kompas():
    d, codes = _codes(os.path.join(REAL, "ДТМЛ.711111.001 - Кольцо.pdf"))
    assert _find(d, "H04", "Ц6.хр")
    d, codes = _codes(os.path.join(REAL, "ДТМЛ.741124.068 - Пластина.pdf"), cfg=_mfg_cfg())
    t01 = [f for f in d.findings if f.code == "T01"]
    assert t01 and t01[0].severity == "info" and "М16" in t01[0].message   # ⌀17 — под крепёж ГОСТ 11284


def test_special_threads_parse():
    p = parse_dimension("Tr40×7-7H")
    assert (p["kind"], p["nominal"], p["pitch"]) == ("thread_tr", 40, 7)
    assert parse_dimension("Tr20x8(P4)")["pitch"] == 4
    assert parse_dimension("G1 1/2-A")["g_size"] == "1 1/2"
    p = parse_dimension("М6-7Н")
    assert p["thread_tol"] == "7H" and p["cyrillic_in_field"]


def _synthetic(tt=(), lines=(), material=""):
    from checking.model import Drawing, PageInfo, TextLine
    d = Drawing("x.pdf", "x.pdf", kind="pdf")
    d.pages = [PageInfo(0, 420, 297)]
    y = 40
    for t in lines:
        d.lines.append(TextLine(t, 0, (60, y, 60 + 3 * len(t), y + 5), (1, 0), 5, (60, y + 5)))
        y += 12
    d.tt = list(tt)
    d.title.material = material
    from checking.material import parse_material as pm
    d.facts["material"] = pm(material)
    return d


def test_text_rules_synthetic():
    from checking.config import load_config
    from checking.textreq import check_tt_text, check_heat, check_refs
    from checking.matcheck import check_material
    from checking.elements import check_elements
    cfg = load_config()
    d = _synthetic(tt=["1. Деталь закалить. HRC 45", "3. Цементировать, HRC 58...62", "4. Общие допуски по ГОСТ 30893.1-x"],
                   lines=["Отв. центр. B3,15 ГОСТ 14034-74", "d-8×36×40H7/f7×6"])
    for fn in (check_tt_text, check_heat, check_refs, check_elements):
        fn(d, cfg)
    codes = {f.code for f in d.findings}
    assert {"P01", "H01", "H06", "P05", "K03"} <= codes and "K01" not in codes
    d = _synthetic(material="Круг 40 ГОСТ 19903-74 / 40Х ГОСТ 1050-88")
    check_material(d, cfg)
    check_refs(d, cfg)
    codes = [f.code for f in d.findings]
    assert "M03" in codes and "M02" in codes and codes.count("N01") == 2


def test_mfg_module_switch():
    from checking import catalog
    assert catalog.rule_settings({}, "T01") == (False, None)                       # блок выключен по умолчанию
    assert catalog.rule_settings({"modules": {"mfg": True}}, "T01")[0] is True
    assert catalog.rule_settings({"modules": {"mfg": True}, "rules": {"T01": {"enabled": False}}}, "T01")[0] is False
    cfg = {"rules": {"T01": {"severity": "error"}}}
    t01 = next(r for g in catalog.catalog_view(cfg) for r in g["rules"] if r["code"] == "T01")
    assert t01["enabled"] is True and t01["severity"] == "error"                   # панель видит собственную настройку
    assert catalog.rule_settings({}, "D03")[0] is True                              # ЕСКД блоком не затронуто


def test_catalog_profile():
    from checking import catalog
    cfg = {"rules": {"D11": {"enabled": False}, "D03": {"severity": "warning"}}, "check_title_people": True}
    assert catalog.rule_settings(cfg, "D11") == (False, None)
    assert catalog.rule_settings(cfg, "D03") == (True, "warning")
    assert catalog.rule_settings(cfg, "B09")[0] is True             # старый ключ config.json
    assert catalog.rule_settings({}, "D23")[0] is False              # по умолчанию выключено
    assert catalog.rule_settings({"rules": {"F00": {"enabled": False}}}, "F00")[0] is True   # служебное
    d = {"findings": [{"code": "D11", "severity": "info", "message": "a", "category": "x"},
                      {"code": "D03", "severity": "error", "message": "b", "category": "x"}], "suppressed": []}
    out = catalog.apply_to_dict(d, cfg)
    assert [f["code"] for f in out["findings"]] == ["D03"] and out["counts"]["warning"] == 1
    assert {r[0] for r in catalog.RULES} >= {"G01", "V03", "N01", "H04", "M02", "Z03", "K01", "D24"}


def _web(tmp_path, monkeypatch, login="admin", password="admin"):
    """Тестовый клиент веб-приложения с изолированными users/feedback/config/work; вход под login."""
    import app as webapp
    from checking import admin, config, feedback, users
    monkeypatch.setattr(config, "CONFIG_PATH", str(tmp_path / "config.json"))
    monkeypatch.setattr(admin, "PATH", str(tmp_path / "admin.json"))
    monkeypatch.setattr(users, "PATH", str(tmp_path / "users.json"))
    monkeypatch.setattr(feedback, "DIR", str(tmp_path / "feedback"))
    monkeypatch.setattr(webapp, "WORK", str(tmp_path / "work"))
    webapp.JOBS.clear()
    c = webapp.app.test_client()
    if login:
        r = c.post("/api/login", json={"name": login, "password": password})
        assert r.status_code == 200, r.json
    return webapp, c


def test_admin_api(tmp_path, monkeypatch):
    from checking import config
    webapp, c = _web(tmp_path, monkeypatch)
    st = c.get("/api/admin/state").json
    assert st["auth"]["user"] == "admin" and any(g["key"] == "gdt" for g in st["groups"])
    r = c.post("/api/admin/settings", json={"rules": {"D11": {"enabled": False}, "D23": {"enabled": True},
                                                      "D03": {"enabled": True, "severity": "warning"}},
                                            "params": {"tight_it": 4}})
    assert r.status_code == 200
    saved = config.load_config()
    assert saved["rules"] == {"D11": {"enabled": False}, "D23": {"enabled": True}, "D03": {"severity": "warning"}}
    assert saved["tight_it"] == 4
    assert c.post("/api/admin/rule", json={"code": "D11", "enabled": True}).status_code == 200
    assert "D11" not in config.load_config()["rules"]
    anon = webapp.app.test_client()
    assert anon.get("/api/admin/state").status_code == 401
    assert anon.post("/api/admin/rule", json={"code": "D11", "enabled": False}).status_code == 401
    assert anon.get("/").status_code == 302 and "/login" in anon.get("/").headers["Location"]
    assert anon.post("/api/login", json={"name": "admin", "password": "bad"}).status_code == 403
    assert anon.post("/api/login", json={"name": "Admin", "password": "admin"}).status_code == 200   # имя без учёта регистра
    assert anon.get("/api/admin/state").status_code == 200


def test_every_rule_code_in_catalog():
    """Каждый код замечания, который выдаёт программа, есть в каталоге — его можно выключить в /admin."""
    import glob
    import re as _re
    from checking import catalog
    used = set()
    for f in glob.glob(os.path.join(ROOT, "checking", "*.py")):
        used |= set(_re.findall(r'add\(\s*(?:ERROR|WARNING|INFO|sev|[a-z_]+)\s*,\s*"([A-Z]\d{2})"', open(f, encoding="utf-8").read()))
    missing = sorted(used - set(catalog.RULE_INDEX))
    assert not missing, missing
    assert len(used) > 100


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "assembly")), reason="нет сборочных чертежей")
def test_assembly_drawings():
    ASM = os.path.join(SAMPLES, "assembly")
    pack = check_file(os.path.join(ASM, "ОН-0020.000 СБ Полный пакет.pdf"))
    opora = pack[2]
    t = opora.facts["tables"][0]
    assert t["kind"] == "spec" and [i["pos"] for i in t["items"] if i["pos"]] == [1, 2, 3, 4, 5, 6]
    assert _find(opora, "A08", "Позиция 6")                          # у поз. 6 не указано количество
    assert not [f for f in opora.findings if f.code in ("A01", "A02")]   # позиции 1–6 на чертеже есть
    assert not [x for x in opora.tt if "Балка (1:4)" in x or "40х40х2" in x]   # таблица и надписи видов — не ТТ
    sb = pack[0]
    assert sorted({p["n"] for p in sb.facts["positions"]}) == [1, 2, 5, 8, 9, 10, 11, 14]
    assert not [f for f in sb.findings if f.code in ("A01", "A02", "A07")]   # резервные номера допустимы
    kol, = check_file(os.path.join(ASM, "Р 38.02 СБ Колонна.pdf"))
    assert not [f for f in kol.findings if f.code == "B01"]          # масса «см. табл.»
    assert kol.tt[0].startswith("1. Конструкция сварная. Конструктивные элементы")
    d = _red("120-001423 Водило 3 ступени в сборе с Иг пошш.pdf")
    assert sorted(p["n"] for p in d.facts["positions"]) == list(range(1, 10))
    assert _find(d, "A05", "нет размеров") and _find(d, "A11", "А4")
    d = _red("120-001427 Водило 2 ступени в сборе с Иг пошш.pdf")
    assert _find(d, "A01", ": 6")


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "assembly")), reason="нет сборочных чертежей")
def test_assembly_mass_across_documents():
    from checking.engine import check_job
    pack = check_file(os.path.join(SAMPLES, "assembly", "ОН-0020.000 СБ Полный пакет.pdf"), apply_profile=False)
    sb, opora, platik = pack[0], pack[2], pack[3]
    check_job(pack, apply_profile=False)
    # реальная ошибка пакета: ОН-0020.011 — «Профиль … L = 650 мм» в Опоре и «Платик … 60х60 мм» в Платике в сборе
    assert _find(platik, "A17", "Одно обозначение присвоено разным изделиям")
    assert sb.facts["mass_parts"] == {"known": 162.64, "parts": 3, "unknown": 5}   # 7×10 + 12×2,38 + 24×2,67
    assert not [f for f in sb.findings if f.code in ("A13", "A15")]                # остальное — крепёж
    assert opora.facts["mass_parts"]["known"] == 7.77                              # «Метал наплавленный 1 | кг»

    pack = check_file(os.path.join(SAMPLES, "assembly", "ОН-0020.000 СБ Полный пакет.pdf"), apply_profile=False)
    sb, opora = pack[0], pack[2]
    sb.title.mass, sb.title.mass_raw = 100.0, "100"
    items = opora.facts["tables"][0]["items"]
    items[5]["qty"] = "1"                                   # поз. 6 с количеством — массы всех частей известны
    items[0]["note"] = "0,5 кг"                             # поз. 1 ОН-0020.010: в чертеже 0,09 кг
    opora.title.mass, opora.title.mass_raw = 14.0, "14"
    changed = check_job(pack, apply_profile=False)
    assert {0, 2} <= changed
    assert _find(sb, "A15", "меньше суммы масс составных частей — 190,64 кг")   # 7×14 + 28,56 + 64,08
    assert _find(opora, "A16", "масса в спецификации 0,5 кг, в чертеже ОН-0020.010 СБ — 0,09 кг")
    assert _find(opora, "A13", "сумма масс всех составных частей 10,01 кг")


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "reductor")), reason="нет чертежей редуктора")
def test_mass_by_view_size():
    d = _red("120-001427 Водило 2 ступени в сборе с Иг пошш.pdf")
    assert _find(d, "B15", "122379,23 кг больше массы сплошного стального бруса")
    for name in ("120-001425 Водило 1 ступени в сборе с Иг пошш.pdf", "280-000036 Водило 1 ступени версия 2.pdf",
                 "280-000092 Водило 3 ступени версия 4.pdf"):
        d = _red(name)
        assert d.facts["mass_view"] and not [f for f in d.findings if f.code == "B15"], name
    for name in ("ДТМЛ.333333.001 СБ - Кронштейн изм.1.pdf",):
        d, = check_file(os.path.join(SAMPLES, "assembly", name))
        assert d.facts["mass_view"]["bound"] > 10 * d.title.mass          # запас ~13 раз


# --- спецификация по ГОСТ 2.106 -----------------------------------------------------------------------
def test_spec_document_kompas_report():
    from checking.specdoc import check_spec_document, children_for
    spec = Spec(os.path.join(SAMPLES, "СП Опора пассивная.xls"))
    d = check_spec_document(spec)
    assert d.kind == "spec" and d.title.designation == "ДТМЛ.301321.001"
    assert _find(d, "S02", "Гайка М18-6H.5.016 ГОСТ 5915-70 — 12 строк")      # повтор вместо общего количества
    assert _find(d, "S01", "Каток") and _find(d, "S03", "Детали")
    assert not [f for f in d.findings if f.code == "S12"]                    # массы сходятся (подшипники учтены)
    ch = children_for(spec, "ДТМЛ.301563.001 СБ")
    assert [c["name"] for c in ch] == ["Труба", "Пластина"]


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "spec")), reason="нет выгрузок СП")
def test_spec_document_erp_and_solid():
    from checking.specdoc import check_spec_document
    d = check_spec_document(Spec(os.path.join(SAMPLES, "spec", "СП_Опора_пассивная_ERP.xlsx")))
    assert _find(d, "S01", "Опора пассивная") and _find(d, "S04", "Шпилька")
    assert not [f for f in d.findings if f.code == "S12"]   # детали с краской — не сборки; у сборок массы сходятся
    d = check_spec_document(Spec(os.path.join(SAMPLES, "spec", "СП_F001-13_Solid.xlsx")))
    s06 = [f for f in d.findings if f.code == "S06"]
    assert s06 and "артикул" in s06[0].message and len(s06) < 20       # одно замечание на сборку, а не на строку


def test_spec_rules_on_drawing_tables():
    d = _red("120-001423 Водило 3 ступени в сборе с Иг пошш.pdf")
    assert _find(d, "S04", "Шайба") and _find(d, "S06", "артикул")
    pack = check_file(os.path.join(SAMPLES, "assembly", "ОН-0020.000 СБ Полный пакет.pdf"))
    assert _find(pack[3], "A14", "ОН-0020.010 СБ")


# --- правила предприятия --------------------------------------------------------------------------------
def test_custom_rules_evaluate():
    from checking import custom
    from checking.report import drawing_to_dict
    d = drawing_to_dict(check_file(os.path.join(REAL, "ДТМЛ.741124.068 - Пластина.pdf"))[0])
    rules = [
        {"code": "U01", "title": "Марки", "kind": "allowed", "what": "grade", "values": ["09Г2С", "45"],
         "severity": "warning", "applies": "parts"},
        {"code": "U02", "title": "Кромки", "kind": "text_required", "where": "tt", "match": "contains",
         "pattern": "Острые кромки притупить", "severity": "info"},
        {"code": "U03", "title": "Организация", "kind": "field", "field": "organization", "check": "contains",
         "value": "Беретта", "severity": "warning"},
        {"code": "U04", "title": "Лист", "kind": "limit", "what": "thickness_max", "value": "16", "severity": "error",
         "message": "Лист {value} мм толще 16"},
        {"code": "U05", "title": "Только для СП", "kind": "text_forbidden", "pattern": "Пластина", "severity": "info",
         "applies": "spec"},
        {"code": "U06", "title": "Сборки", "kind": "text_required", "pattern": "xyz", "severity": "info",
         "applies": "parts", "material_contains": "09Г2С"},
    ]
    for r in rules:
        assert not custom.validate(r), (r["code"], custom.validate(r))
    hits = {r["code"]: custom.evaluate(d, r) for r in rules}
    assert "Ст3сп" in hits["U01"][0]["message"]                     # марки нет в перечне
    assert hits["U02"] and not hits["U03"]                           # в ТТ нет «Острые кромки…»; ООО «Беретта» есть
    assert hits["U04"][0]["message"] == "Лист 20 мм толще 16" and hits["U04"][0]["severity"] == "error"
    assert not hits["U05"] and not hits["U06"]                       # не тот тип документа / материал
    assert custom.validate({"kind": "text_required", "title": "x", "severity": "info", "match": "regex",
                            "pattern": "("})                         # битое регулярное выражение
    cfg = {"custom_rules": rules, "rules": {"U02": {"enabled": False}}}
    from checking import catalog
    out = catalog.apply_to_dict(custom.apply_to_dict(d, cfg), cfg)
    codes = {f["code"] for f in out["findings"]}
    assert {"U01", "U04"} <= codes and "U02" not in codes            # U02 выключено в списке правил


def test_custom_rules_api(tmp_path, monkeypatch):
    from checking import config
    webapp, c = _web(tmp_path, monkeypatch)
    bad = c.post("/api/admin/custom", json={"title": "", "kind": "allowed", "severity": "info"})
    assert bad.status_code == 400 and "Перечень пуст" in bad.json["error"]
    r = c.post("/api/admin/custom", json={"title": "Марки", "kind": "allowed", "what": "grade", "severity": "warning",
                                          "values": "Ст3сп\n09Г2С"})
    assert r.status_code == 200 and r.json["code"] == "U01"
    assert any(g["key"] == "custom" and g["rules"][0]["code"] == "U01" for g in r.json["groups"])
    assert c.post("/api/admin/rule", json={"code": "U01", "enabled": False}).status_code == 200
    assert config.load_config()["rules"]["U01"] == {"enabled": False}
    r = c.post("/api/admin/custom", json={"code": "U01", "title": "Марки 2", "kind": "allowed", "what": "grade",
                                          "severity": "info", "values": ["45"]})
    assert r.json["code"] == "U01" and config.load_config()["custom_rules"][0]["title"] == "Марки 2"
    assert c.delete("/api/admin/custom/U01").json["custom_rules"] == []
    assert "U01" not in (config.load_config().get("rules") or {})


def test_graphic_roughness_sign_not_reported_as_missing():
    d, = check_file(os.path.join(SAMPLES, "drawings", "ДТМЛ.715311.003 - Палец нижний.pdf"))
    assert d.facts["rough_signs"] >= 1                       # знак «без обработки» нарисован графикой, без текста
    assert not [f for f in d.findings if f.code == "R02"]
    d, = check_file(os.path.join(SAMPLES, "drawings", "ДТМЛ.711111.001 - Кольцо.pdf"))
    assert [f for f in d.findings if f.code == "R02"]        # здесь знаков на изображении действительно нет


def test_dimension_layout(tmp_path):
    import pymupdf
    k = 72 / 25.4
    doc = pymupdf.open()
    pg = doc.new_page(width=297 * k, height=210 * k)

    def line(x1, y1, x2, y2):
        pg.draw_line((x1 * k, y1 * k), (x2 * k, y2 * k), width=0.5)

    def text(x, y, t, size=14):
        pg.insert_text((x * k, y * k), t, fontsize=size, fontname="helv")
    pg.draw_rect(pymupdf.Rect(20 * k, 5 * k, 292 * k, 205 * k), width=1)
    pg.draw_rect(pymupdf.Rect(60 * k, 60 * k, 160 * k, 120 * k), width=1)
    line(60, 58, 60, 52); line(160, 58, 160, 52); line(60, 55, 160, 55); text(100, 53.5, "100")   # 5 мм от контура
    line(60, 122, 60, 146); line(160, 122, 160, 134); line(60, 132, 160, 132); text(100, 130.5, "100")
    line(110, 122, 110, 146); line(60, 136, 110, 136); text(75, 134.5, "50")                      # 4 мм от соседней
    line(58, 60, 42, 60); line(58, 120, 42, 120); line(45, 60, 45, 120)                            # 15 мм — норма
    pg.insert_text((43 * k, 95 * k), "60", fontsize=14, fontname="helv", rotate=90)
    text(200, 100, "Size 45 mm"); text(203, 103, "Tolerance 7 um")                                 # наложение
    text(30, 180, "Технические требования проверка расположения размеров на листе чертежа", 10)
    p = tmp_path / "layout.pdf"
    doc.save(str(p))
    d, = check_file(str(p))
    got = {(f.code, f.message.split("»")[0]) for f in d.findings if f.code in ("D31", "D32", "D33")}
    assert any(c == "D32" and "«100" in m for c, m in got)
    assert [f for f in d.findings if f.code == "D33" and "«100» и «50»" in f.message]
    assert [f for f in d.findings if f.code == "D31"]
    assert not [f for f in d.findings if f.code == "D32" and "«60»" in f.message]


@pytest.mark.skipif(not os.path.isdir(RED), reason="нет образцов редуктора")
def test_vertical_dimension_tolerances_not_stolen_by_neighbour():
    d = _red("280-000060 Зубчатый венец редуктора Прево 2 ступени.pdf")
    raws = [x.raw for x in d.dims]
    assert raws.count("⌀335 -0,1 / -0,2") == 2 and raws.count("⌀295,7 +0,1") == 2
    assert not [r for r in raws if r.startswith("⌀295,7") and "-0,1" in r]


def test_users_roles_profile_and_feedback(tmp_path, monkeypatch):
    import json as _json
    from checking import feedback
    webapp, adm = _web(tmp_path, monkeypatch)
    assert adm.get("/api/me").json["default_password"] is True              # admin/admin — пароль по умолчанию
    assert adm.post("/api/admin/users", json={"name": "ivanov", "password": "pw1234"}).status_code == 200
    assert adm.post("/api/admin/users", json={"name": "ivanov", "password": "x"}).status_code == 400
    u = webapp.app.test_client()
    assert u.post("/api/login", json={"name": "ivanov", "password": "pw1234"}).status_code == 200
    assert u.get("/api/me").json["role"] == "user"
    # обычному пользователю настройки и журнал закрыты
    assert u.get("/api/admin/state").status_code == 403 and u.get("/admin").status_code == 403
    assert u.get("/api/admin/actions").status_code == 403 and u.post("/api/exceptions", json={}).status_code == 403
    # смена пароля в профиле
    assert u.post("/api/profile/password", json={"old": "bad", "new": "newpass"}).status_code == 403
    assert u.post("/api/profile/password", json={"old": "pw1234", "new": "newpass"}).status_code == 200
    assert webapp.app.test_client().post("/api/login", json={"name": "ivanov", "password": "pw1234"}).status_code == 403
    assert u.post("/api/login", json={"name": "ivanov", "password": "newpass"}).status_code == 200
    # проверка принадлежит пользователю; другой пользователь её не видит, администратор видит
    jdir = tmp_path / "work" / "20260101-000000-aaaaaa"
    jdir.mkdir(parents=True)
    rep = {"file": "a.pdf", "id": 0, "title": {"designation": "X.1", "name": "n"}, "counts": {"error": 0, "warning": 1, "info": 0},
           "findings": [{"severity": "warning", "code": "D10", "category": "c", "message": "m", "page": 0}], "pages": [], "dims": [],
           "tt": [], "roughness": [], "limits": [], "facts": {}, "suppressed": []}
    (jdir / "report.json").write_text(_json.dumps([rep]), encoding="utf-8")
    (jdir / "meta.json").write_text(_json.dumps({"owner": "ivanov"}), encoding="utf-8")
    jid = "20260101-000000-aaaaaa"
    assert u.get(f"/api/jobs/{jid}").status_code == 200 and adm.get(f"/api/jobs/{jid}").status_code == 200
    assert adm.post("/api/admin/users", json={"name": "petrov", "password": "pw1234"}).status_code == 200
    p = webapp.app.test_client()
    p.post("/api/login", json={"name": "petrov", "password": "pw1234"})
    assert p.get(f"/api/jobs/{jid}").status_code == 403 and p.get("/api/history").json == []
    assert [h["id"] for h in u.get("/api/history").json] == [jid] and adm.get("/api/history").json[0]["owner"] == "ivanov"
    # отметки: верно/ложное, пропуск, ошибка разбора
    f = {"type": "finding", "verdict": "false", "doc": 0, "code": "D10", "message": "m", "comment": "общие допуски в ТТ"}
    assert u.post(f"/api/jobs/{jid}/feedback", json=f).status_code == 200
    assert u.post(f"/api/jobs/{jid}/feedback", json={**f, "verdict": "ok"}).status_code == 200      # перезаписывает свою отметку
    assert u.post(f"/api/jobs/{jid}/feedback", json={"type": "missed", "doc": 0, "page": 0, "point": [10, 20],
                                                      "group": "Размеры", "comment": ""}).status_code == 400
    assert u.post(f"/api/jobs/{jid}/feedback", json={"type": "missed", "doc": 0, "page": 0, "point": [10, 20],
                                                      "group": "Размеры", "comment": "нет шероховатости"}).status_code == 200
    assert u.post(f"/api/jobs/{jid}/feedback", json={"type": "parse", "doc": 0, "dim": "⌀295,7 +0,1 / -0,1",
                                                      "expected": "⌀295,7 +0,1"}).status_code == 400   # нужен комментарий
    mine = u.get(f"/api/jobs/{jid}/feedback").json
    assert [e["type"] for e in mine] == ["finding", "missed"] and mine[0]["verdict"] == "ok"
    # комментарий к замечанию — только администратор; пользователь его не видит
    cm = {"type": "comment", "doc": 0, "code": "D10", "message": "m", "comment": "размеры 12 и 15 стоят рядом",
          "expected": "два размера: 12 и 15", "page": 0, "bbox": [1, 2, 3, 4]}
    assert u.post(f"/api/jobs/{jid}/feedback", json=cm).status_code == 403
    assert adm.post(f"/api/jobs/{jid}/feedback", json={**cm, "comment": ""}).status_code == 400
    r = adm.post(f"/api/jobs/{jid}/feedback", json=cm)
    assert r.status_code == 200 and r.json["type"] == "comment"
    assert [e["type"] for e in u.get(f"/api/jobs/{jid}/feedback").json] == ["finding", "missed"]
    assert [e["comment"] for e in adm.get(f"/api/jobs/{jid}/feedback").json if e["type"] == "comment"] == [cm["comment"]]
    sm0 = adm.get("/api/admin/feedback").json
    assert sm0["comments"][0]["expected"] == cm["expected"] and not [e for e in sm0["parse"] if e["type"] == "comment"]
    assert adm.delete(f"/api/jobs/{jid}/feedback/{r.json['id']}").status_code == 200
    sm = adm.get("/api/admin/feedback").json
    assert sm["rules"][0]["code"] == "D10" and sm["rules"][0]["ok"] == 1 and sm["missed_groups"] == [["Размеры", 1]]
    acts = adm.get("/api/admin/actions").json
    assert {"Вход", "Смена пароля", "Отметка замечания", "Отмечен пропуск"} <= {a["action"] for a in acts}
    assert all(a["user"] == "ivanov" for a in u.get("/api/profile/activity").json)
    # выгрузка для разработки
    import io, zipfile
    z = zipfile.ZipFile(io.BytesIO(adm.get("/api/admin/export.zip").data))
    names = z.namelist()
    assert "actions.jsonl" in names and f"jobs/{jid}/reviews.json" in names and f"jobs/{jid}/report.json" in names
    assert "hash" not in z.read("users.json").decode()
    # нельзя удалить последнего администратора и себя
    assert adm.delete("/api/admin/users/admin").status_code == 400
    assert adm.patch("/api/admin/users/admin", json={"role": "user"}).status_code == 400
    assert adm.delete("/api/admin/users/petrov").status_code == 200
    assert p.get("/api/me").json is None or p.get("/api/history").status_code == 401           # удалённый пользователь теряет доступ

def test_custom_rule_when_tt_and_thread_tol():
    from checking import custom
    ost = r"ОСТ\s*4\s*Г[О0O]\.?\s*070\.?\s*014"
    base = {"title": "x", "severity": "warning", "kind": "allowed", "what": "thread_tol", "values": ["6g", "6H"], "when_tt": ost}
    assert not custom.validate(base)
    assert custom.validate({**base, "when_tt": "("})
    d = {"title": {}, "facts": {}, "tt": ["Технические требования по ОСТ4 Г0.070.014"], "texts": [],
         "dims": [{"kind": "thread", "raw": "M8-7H", "thread_tol": "7H"}, {"kind": "thread", "raw": "M6-6H", "thread_tol": "6H"}]}
    hits = custom.evaluate(d, {**base, "code": "U90"})
    assert len(hits) == 1 and "7H" in hits[0]["message"]
    d["tt"] = ["Острые кромки притупить"]                      # ссылки на ОСТ нет — правило не действует
    assert not custom.evaluate(d, {**base, "code": "U90"})


def test_obsolete_30893():
    from checking.normdata import OBSOLETE_STANDARDS
    assert ("ГОСТ 25670-83", "ГОСТ 30893.1-2002") in OBSOLETE_STANDARDS

def test_ost_070_014():
    from checking.model import Drawing, Dimension
    from checking.ost import check_ost
    from checking import catalog
    drw = Drawing(path="x.pdf", filename="x.pdf")
    drw.tt = ["1. Направление проката не указывать", "2. Остальные технические требования по ОСТ4 Г0.070.014"]
    drw.dims = [Dimension(raw="M8-7H", page=0, bbox=(0, 0, 1, 1), kind="thread", thread_tol="7H"),
                Dimension(raw="M6-6g", page=0, bbox=(0, 0, 1, 1), kind="thread", thread_tol="6g")]
    check_ost(drw, {})
    codes = [f.code for f in drw.findings]
    o01 = next(f for f in drw.findings if f.code == "O01")
    assert codes.count("O01") == 1 and "7H" in o01.message and "O02" in codes
    drw2 = Drawing(path="y.pdf", filename="y.pdf")
    drw2.tt = ["Направление проката не указывать"]
    drw2.dims = list(drw.dims)
    check_ost(drw2, {})
    assert not drw2.findings                                        # нет ссылки на ОСТ — требования не действуют
    view = {g["key"]: g for g in catalog.catalog_view({})}
    assert {r["code"] for r in view["ost070014"]["rules"]} >= {"O01", "O02"} and view["ost070014"]["clauses"]
    assert all(g["section"] for g in view.values())



def test_rule_standards_fresh():
    """checking/rule_standards.py соответствует исходникам (иначе: python tools/gen_rule_standards.py)."""
    import importlib.util
    spec = importlib.util.spec_from_file_location("gen_rule_standards", os.path.join(ROOT, "tools", "gen_rule_standards.py"))
    gen = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gen)
    from checking.rule_standards import RULE_STANDARDS
    assert gen.collect() == RULE_STANDARDS


def test_standards_index():
    from checking import catalog
    ix = {e["key"]: e for e in catalog.standards_index({})}
    assert ix["ГОСТ 25347"]["label"] == "ГОСТ 25347-2013"
    d = {x["code"]: x["clauses"] for x in ix["ГОСТ Р 2.109"]["direct"]}
    assert d["A01"] == ["п. 7.3.1"] and d["Q19"] == ["п. 9.1"]
    assert {x["code"] for x in ix["ОСТ4 ГО.070.015"]["direct"]} >= {"O03", "O04", "O19", "O26", "O30"}
    assert catalog.parse_standards("ГОСТ Р 2.104-2023, ГОСТ 2.106-96 (2019)")[0][0] == "ГОСТ Р 2.104"


def test_ost_070_015():
    from checking.model import Drawing
    from checking.ost import check_ost
    drw = Drawing(path="x.pdf", filename="x.pdf")
    drw.tt = ["1. Технические требования по ОСТ 4 Г0.070.015", "2. Заклёпки развальцевать"]
    drw.facts["is_assembly"] = True
    check_ost(drw, {})
    assert [f.code for f in drw.findings if f.code != "O19"] == ["O03"]
    drw.tt.append("3. Места клёпки покрыть лаком АК-113")
    drw.findings.clear()
    check_ost(drw, {})
    assert [f.code for f in drw.findings] == ["O19"]                # ссылка не в форме «Остальные технические требования»
    part = Drawing(path="y.pdf", filename="y.pdf")
    part.tt = ["Технические требования по ОСТ 4 ГО.070.015"]
    check_ost(part, {})
    assert [f.code for f in part.findings] == ["O04"]


def _ost_drw(tt, material="", asm=False):
    from checking.model import Drawing
    d = Drawing(path="x.pdf", filename="x.pdf")
    d.tt = list(tt)
    d.title.material = material
    d.facts["is_assembly"] = asm
    return d


def _ost_codes(tt, material="", asm=False):
    from checking.ost import check_ost
    d = _ost_drw(tt, material, asm)
    check_ost(d, {})
    return [f.code for f in d.findings]


def test_ost_070_014_full():
    ref = "Остальные технические требования по ОСТ4 Г0.070.014"
    assert _ost_codes([ref]) == []                                           # чистый чертёж — без замечаний
    assert _ost_codes(["Неуказанные кромки"]) == []                          # ссылки нет — ничего не проверяется
    assert _ost_codes(["Общие допуски по ГОСТ 30893.2-H", ref]) == ["O05"]
    assert _ost_codes(["Общие допуски по ГОСТ 30893.2-К", ref]) == []
    assert _ost_codes(["Неуказанные предельные отклонения: H12, h12, ±IT12/2", ref]) == ["O06"]
    assert _ost_codes(["Неуказанные предельные отклонения: H14, h14, ±IT14/2", ref]) == []
    assert _ost_codes(["Неуказанные отклонения углов ±AT12/2", ref]) == ["O07"]
    assert _ost_codes(["Угловое положение отверстий ±5°", ref]) == ["O18"]
    assert _ost_codes(["Острые кромки притупить R0,5", ref]) == ["O08"]
    assert _ost_codes(["Острые кромки притупить R0,2...0,3", ref]) == []
    assert _ost_codes(["Радиусы галтелей R1 max", ref]) == ["O09"]
    assert _ost_codes(["Радиусы галтелей не более 0,5", ref]) == []
    assert _ost_codes(["Резьба Rz80", ref]) == ["O10"]
    assert _ost_codes(["Направление проката не указывать", ref]) == ["O02"]
    assert "O17" in _ost_codes(["Технические требования по ОСТ4 ГО.070.014 ред. 1-70"])
    assert _ost_codes(["Технические требования по ОСТ4 ГО.070.014", "Покрытие Ан.Окс"]) == ["O16"]
    assert _ost_codes([ref], asm=True) == ["O15"]
    # материалы
    assert _ost_codes([ref], "Фанера") == ["O11"]
    assert _ost_codes([ref], "Фанера ГОСТ 3916.1-96") == []
    assert _ost_codes([ref], "Сосна ГОСТ 2695-83") == ["O11"]               # сосна — хвойные, нужен ГОСТ 8486
    assert _ost_codes([ref], "Берёза ГОСТ 2695-83") == []
    assert _ost_codes(["Неуказанные предельные отклонения: H14, h14, ±IT14/2", ref], "Берёза ГОСТ 2695-83") == ["O12"]
    assert _ost_codes(["Неуказанные предельные отклонения: H14, h14, ±IT14/2", ref], "Резина ТМКЩ") == ["O13"]
    assert _ost_codes(["Допуски по ТУ 2500-376-00152106-94", ref], "Резина ТМКЩ") == []
    assert _ost_codes(["Острые кромки притупить R1", ref], "Стекло") == ["O14", "O08"] or \
        set(_ost_codes(["Острые кромки притупить R1", ref], "Стекло")) == {"O14", "O08"}
    assert _ost_codes(["Острые кромки притупить R1", ref], "Стекло органическое") == ["O08"]


def test_ost_codes_in_catalog():
    from checking import catalog
    codes = {r[0] for r in catalog.RULES if r[1] in ("ost070014", "ost070015")}
    assert codes == {f"O{i:02d}" for i in range(1, 33)}
    cl15 = {c.strip() for _, _, cc in catalog.CLAUSES["ost070015"] if cc for c in cc.split(",")}
    assert cl15 == {r[0] for r in catalog.RULES if r[1] == "ost070015"}          # каждое правило есть в таблице пунктов
    cl = {c for _, _, c in catalog.CLAUSES["ost070014"] if c}
    assert cl == {r[0] for r in catalog.RULES if r[1] == "ost070014"}

def test_rule_codes_backward_compatible():
    """Коды правил из версии, которую тестируют пользователи, не удалены и не изменены (docs/MIGRATION.md)."""
    import json
    from checking import catalog
    base = json.load(open(os.path.join(ROOT, "tests", "rule_codes_baseline.json"), encoding="utf-8"))["rules"]
    problems = []
    for code, b in base.items():
        r = catalog.RULE_INDEX.get(code)
        if r is None:
            problems.append(f"{code}: правила нет в каталоге")
        elif (r[1], bool(r[3]), r[4]) != (b["group"], b["default"], b["severity"]):
            problems.append(f"{code}: было {b}, стало группа={r[1]}, по умолчанию={r[3]}, важность={r[4]}")
    assert not problems, problems


def test_rule_docs_complete():
    from checking import catalog, rule_docs
    codes = {r[0] for r in catalog.RULES}
    assert set(rule_docs.RULE_DOCS) == codes, (sorted(codes - set(rule_docs.RULE_DOCS)), sorted(set(rule_docs.RULE_DOCS) - codes))
    for code, d in rule_docs.RULE_DOCS.items():
        assert d["what"].strip() and d["how"].strip() and d["example"].strip() and d["skip"].strip(), code
    assert set(rule_docs.REV) <= codes
    view = [r for g in catalog.catalog_view({}) for r in g["rules"]]
    assert all(r["doc"] and r["rev"] >= 1 for r in view)


def test_feedback_revision(tmp_path, monkeypatch):
    from checking import feedback, rule_docs
    monkeypatch.setattr(feedback, "DIR", str(tmp_path))
    e = feedback.add("job1", "u", {"type": "finding", "verdict": "false", "code": "D03", "message": "m", "doc": "a"})
    assert e["rev"] == 1 and e["app"]
    # отметка старой версии программы — без поля rev
    old = {"id": "x1", "ts": "2026-01-01 10:00:00", "user": "u", "job": "job2", "type": "finding", "verdict": "false",
           "code": "D03", "message": "m"}
    import json
    os.makedirs(os.path.join(str(tmp_path), "reviews"), exist_ok=True)
    json.dump([old], open(os.path.join(str(tmp_path), "reviews", "job2.json"), "w", encoding="utf-8"))
    monkeypatch.setitem(rule_docs.REV, "D03", 2)               # поведение правила изменили
    row = next(r for r in feedback.summary()["rules"] if r["code"] == "D03")
    assert row["total"] == 0 and row["old"]["total"] == 2 and row["rev"] == 2    # обе отметки — до изменения
    e2 = feedback.add("job3", "u", {"type": "finding", "verdict": "ok", "code": "D03", "message": "m2", "doc": "b"})
    assert e2["rev"] == 2
    row = next(r for r in feedback.summary()["rules"] if r["code"] == "D03")
    assert row["ok"] == 1 and row["false"] == 0 and row["old"]["false"] == 2
    monkeypatch.setitem(rule_docs.REV, "D03", 1)
    unknown = {**old, "id": "x2", "code": "Z99"}
    json.dump([unknown], open(os.path.join(str(tmp_path), "reviews", "job4.json"), "w", encoding="utf-8"))
    assert next(r for r in feedback.summary()["rules"] if r["code"] == "Z99")["known"] is False   # правила нет в версии

def test_import_exceptions(tmp_path, monkeypatch):
    from checking import feedback
    monkeypatch.setattr(feedback, "DIR", str(tmp_path))
    rules = [{"id": "a1", "scope": "drawing", "code": "D10", "message": "m", "designation": "X.1", "created": "2026-09-01 10:00"},
             {"id": "a2", "scope": "global", "code": "D11", "pattern": "Ra40", "comment": "не нужно"},
             {"id": "a3", "scope": "global"}]                           # без кода — пропускается
    assert feedback.import_exceptions(rules) == 2
    assert feedback.import_exceptions(rules) == 0                       # повторный запуск — без дублей
    rows = {r["code"]: r for r in feedback.summary()["rules"]}
    assert rows["D10"]["false"] == 1 and rows["D11"]["false"] == 1

def test_expert_notes(tmp_path, monkeypatch):
    from checking import expert
    path = str(tmp_path / "feedback" / "expert_notes.json")
    monkeypatch.setattr(expert, "PATH", path)
    assert isinstance(expert.add("D03", "   ", "u"), str)                       # пустое мнение
    n = expert.add("D03", "Сверять отклонения и для js", "ivanov", rev=1, app="2026.10")
    assert n["status"] == "new" and n["code"] == "D03"
    assert [x["id"] for x in expert.listing(status="new")] == [n["id"]]
    assert isinstance(expert.reply(n["id"], "applied", "", 2), str)             # без комментария ответить нельзя
    assert isinstance(expert.reply(n["id"], "new", "x"), str)
    r = expert.reply(n["id"], "applied", "Проверено по ГОСТ 25346, правило изменено", rev_after=2)
    assert r["status"] == "applied" and r["dev"]["rev_after"] == 2
    assert not expert.listing(status="new") and expert.by_rule()["D03"][0]["dev"]["comment"]
    assert not expert.delete(n["id"], "ivanov", False)                          # разобранное мнение не удаляется
    n2 = expert.add("D03", "Ещё мнение", "ivanov")
    assert not expert.delete(n2["id"], "petrov", False) and expert.delete(n2["id"], "ivanov", False)


def test_expert_notes_api(tmp_path, monkeypatch):
    from checking import expert
    monkeypatch.setattr(expert, "PATH", str(tmp_path / "feedback" / "expert_notes.json"))
    webapp, adm = _web(tmp_path, monkeypatch)
    assert adm.post("/api/admin/expert-notes", json={"code": "ZZZ", "text": "x"}).status_code == 400   # такого правила нет
    assert adm.post("/api/admin/expert-notes", json={"code": "O01", "text": "  "}).status_code == 400
    ok = adm.post("/api/admin/expert-notes", json={"code": "O01", "text": "Допустить 6e для наружной"})
    assert ok.status_code == 200 and ok.json["rev"] == 1 and ok.json["user"] == "admin"
    got = adm.get("/api/admin/expert-notes").json
    assert got["notes"]["O01"][0]["status"] == "new" and "applied" in got["statuses"]
    assert adm.delete("/api/admin/expert-notes/" + ok.json["id"]).status_code == 200
    # обычный пользователь мнений не видит и не пишет
    assert adm.post("/api/admin/users", json={"name": "ivanov", "password": "pw1234"}).status_code == 200
    u = webapp.app.test_client()
    assert u.post("/api/login", json={"name": "ivanov", "password": "pw1234"}).status_code == 200
    assert u.get("/api/admin/expert-notes").status_code == 403
    assert u.post("/api/admin/expert-notes", json={"code": "O01", "text": "x"}).status_code == 403

def test_rule_review_marks(tmp_path, monkeypatch):
    """Отметка «утверждено» общая для администраторов, снимается любым из них и новым мнением."""
    from checking import expert
    monkeypatch.setattr(expert, "PATH", str(tmp_path / "feedback" / "expert_notes.json"))
    monkeypatch.setattr(expert, "REVIEWS_PATH", str(tmp_path / "feedback" / "rule_reviews.json"))
    webapp, adm = _web(tmp_path, monkeypatch)
    assert adm.post("/api/admin/rule-review", json={"code": "ZZZ", "reviewed": True}).status_code == 400
    r = adm.post("/api/admin/rule-review", json={"code": "O01", "reviewed": True})
    assert r.status_code == 200 and r.json["review"]["user"] == "admin"
    assert adm.get("/api/admin/expert-notes").json["reviews"]["O01"]["rev"] == 1
    assert adm.post("/api/admin/users", json={"name": "admin2", "password": "pw1234", "role": "admin"}).status_code in (200, 400)
    # новое мнение снимает отметку
    assert adm.post("/api/admin/expert-notes", json={"code": "O01", "text": "Пересмотреть"}).status_code == 200
    assert "O01" not in adm.get("/api/admin/expert-notes").json["reviews"]
    adm.post("/api/admin/rule-review", json={"code": "O02", "reviewed": True})
    assert adm.post("/api/admin/rule-review", json={"code": "O02", "reviewed": False}).json["review"] is None
    u_ok = adm.post("/api/admin/users", json={"name": "ivanov", "password": "pw1234"})
    u = webapp.app.test_client()
    u.post("/api/login", json={"name": "ivanov", "password": "pw1234"})
    assert u.post("/api/admin/rule-review", json={"code": "O01", "reviewed": True}).status_code == 403


# ---------------------------------------------------------------- текстовые документы (ГОСТ Р 2.105-2019) ----
def _textdoc_codes(builder, tmp_path):
    sys.path.insert(0, os.path.join(ROOT, "tests"))
    import make_textdoc
    path = str(tmp_path / (builder + ".pdf"))
    getattr(make_textdoc, builder)(path)
    d = check_file(path, apply_profile=False)[0]
    return d, [f.code for f in d.findings]


def _needs_fonts():
    sys.path.insert(0, os.path.join(ROOT, "tests"))
    import make_textdoc
    return pytest.mark.skipif(not make_textdoc.AVAILABLE, reason="нет шрифтов Times New Roman / Arial")


@_needs_fonts()
def test_textdoc_good_document(tmp_path):
    d, codes = _textdoc_codes("good_document", tmp_path)
    assert d.facts.get("text_document") is True
    bad = [f for f in d.findings if f.severity in ("error", "warning")]
    assert not bad, [(f.code, f.message) for f in bad]
    assert not any(c.startswith(("D", "T", "B0", "E")) for c in codes)            # проверки чертежа не применяются


@_needs_fonts()
def test_textdoc_bad_document(tmp_path):
    d, codes = _textdoc_codes("bad_document", tmp_path)
    expected = {"L08", "L09", "L10", "L13", "L14", "L15", "L16", "L17", "L18", "L19", "L21", "L23", "L25", "L26", "L28",
                "L29", "L31", "L32", "L33", "L34", "L44", "L45", "L46", "L47", "L48", "L49", "L50", "L51", "L52", "L53",
                "L55", "L56", "L57", "L58", "L59", "L60", "L63"}
    assert expected <= set(codes), sorted(expected - set(codes))
    assert not any(c.startswith(("D0", "D1", "T0")) for c in codes)


@_needs_fonts()
def test_textdoc_bad_document2(tmp_path):
    d, codes = _textdoc_codes("bad_document2", tmp_path)
    expected = {"L04", "L05", "L07", "L15", "L17", "L20", "L22", "L24", "L27", "L30", "L35", "L61"}
    assert expected <= set(codes), sorted(expected - set(codes))
    assert "L28" not in codes                                                  # «таблицу 1 и ГОСТ …» — не ссылка на пункт ГОСТ


@_needs_fonts()
def test_textdoc_typography_and_size(tmp_path):
    d, codes = _textdoc_codes("bad_typography", tmp_path)
    assert {"L36", "L38", "L40", "L43"} <= set(codes), codes
    d, codes = _textdoc_codes("long_document", tmp_path)
    assert "L03" in codes                                                      # 27 листов без содержания


def test_textdoc_detection_rules():
    from checking.model import Drawing, PageInfo, TextLine
    from checking.textdoc import is_text_document
    d = Drawing("x.pdf", "x.pdf")
    d.pages = [PageInfo(0, 210, 297)]
    assert is_text_document(d, "ДТМЛ.123456.001ТУ", {}) is True
    assert is_text_document(d, "ДТМЛ.123456.001 ПЗ", {}) is True
    assert is_text_document(d, "ДТМЛ.123456.001СБ", {}) is False              # сборочный чертёж
    assert is_text_document(d, "ДТМЛ.123456.001", {}) is False
    d.pages = [PageInfo(0, 297, 210)]
    assert is_text_document(d, "ДТМЛ.123456.001ТУ", {}) is False              # альбомный лист — не текстовый документ
    d.pages = [PageInfo(0, 210, 297)]
    assert is_text_document(d, "ДТМЛ.123456.001РЭ", {"text_doc_codes": ["ТУ"]}) is False   # список кодов настраивается


def test_textdoc_codes_in_catalog():
    import re as _re
    from checking import catalog
    src = open(os.path.join(ROOT, "checking", "textdoc.py"), encoding="utf-8").read()
    used = set(_re.findall(r'"(L\d\d)"', src))
    in_cat = {r[0] for r in catalog.RULES if r[1] == "textdoc"}
    assert used == in_cat, (sorted(used ^ in_cat))
    assert len(in_cat) == 63
    idx = {e["key"]: e for e in catalog.standards_index({})}
    assert len(idx["ГОСТ Р 2.105"]["direct"]) >= 60

# ------------------------------------------------ правила по основным стандартам ЕСКД (eskd.py) ----
def _eskd_drw(tt=None, lines=(), pages=((210, 297),), **title):
    from checking.model import Drawing, PageInfo, TextLine
    d = Drawing("x.pdf", "x.pdf")
    d.pages = [PageInfo(i, w, h) for i, (w, h) in enumerate(pages)]
    d.title.found = True
    d.title.origin = (20, 240)
    for k, v in title.items():
        setattr(d.title, k, v)
    d.tt = list(tt or [])
    d.facts["_tt_items"] = [(t, 0, None) for t in d.tt]
    for i, (text, x, y, size) in enumerate(lines):
        d.lines.append(TextLine(text, 0, (x, y, x + 3 * len(text) * size / 5, y + size), (1, 0), size, (x, y + size)))
    return d


def _eskd_codes(d, fn=None):
    from checking import eskd
    eskd.check_eskd(d, {})
    return [f.code for f in d.findings]


def test_eskd_formats_scale_designation():
    d = _eskd_drw(format="А4", scale="1:2", designation="ДТМЛ.123456.001СБ")
    assert not [c for c in _eskd_codes(d) if c.startswith("Q")]
    d = _eskd_drw(pages=((215, 300),), format="A4", scale="М 1:2", designation="ДТМ.123456.001")
    codes = _eskd_codes(d)
    assert {"Q01", "Q03", "Q04", "Q41"} <= set(codes) and "Q02" not in codes
    assert "Q02" in _eskd_codes(_eskd_drw(format="А4×2"))
    assert "Q01" not in _eskd_codes(_eskd_drw(pages=((297, 630),)))                     # А4×3
    assert "Q01" not in _eskd_codes(_eskd_drw(pages=((420, 297),)))                     # А3 альбомный
    assert "Q06" in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001-00"))
    assert not any(c.startswith("Q0") and c != "Q03" for c in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001-1")))   # ГОСТ Р 2.201-2023 п. 5.4.7
    assert not any(c in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.000")) for c in ("Q04", "Q06"))
    assert not any(c in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.0012")) for c in ("Q04", "Q06"))   # п. 5.2.5: знаков больше трёх
    assert not any(c in _eskd_codes(_eskd_drw(designation="12345678.123456.001")) for c in ("Q04", "Q05"))  # п. 5.2.3: цифровой код
    assert "Q04" in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.01"))
    assert "Q05" in _eskd_codes(_eskd_drw(designation="DTML.123456.001"))                                   # латиница в коде организации
    assert "Q05" not in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001-A"))                             # п. 4.6
    assert "Q07" in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001СБЧЕРТ"))
    assert "Q05" in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001CБ"))
    assert not any(c in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001-02СБ")) for c in ("Q04", "Q05", "Q06", "Q07"))
    assert "Q04" not in _eskd_codes(_eskd_drw(designation="123-456789"))               # другая система обозначений
    # обозначение на каждом листе
    d = _eskd_drw(pages=((210, 297), (210, 297)), designation="ДТМЛ.123456.001",
                  lines=[("ДТМЛ.123456.001", 100, 280, 3.5)])
    d.lines[0].page = 0
    assert "Q08" not in _eskd_codes(d)
    from checking.model import Drawing, TextLine
    d.lines.append(TextLine("Лист 2", 1, (10, 10, 30, 14), (1, 0), 3.5, (10, 14)))
    d.findings.clear()
    assert "Q08" in _eskd_codes(d)


def test_eskd_references_and_tt():
    ok = ["1. Резьба по ГОСТ 24705", "2. Покрытие по ТУ 123-45", "3. Шероховатость по ГОСТ 2789"]
    assert not [c for c in _eskd_codes(_eskd_drw(tt=ok)) if c in ("Q09", "Q10", "Q11", "Q12", "Q20", "Q21")]
    codes = _eskd_codes(_eskd_drw(tt=["1. Швы по ГОСТ 5264-80 п. 4.2", "2. Изготовить по ТУ", "3. Канавку по ГОСТ 8820",
                                    "4. Сверлить отверстие ⌀5", "5. Температура -40 °С", "6. Длина 25мм"]))
    assert {"Q09", "Q10", "Q11", "Q12", "Q20", "Q21", "Q22"} <= set(codes), codes
    assert "Q12" not in _eskd_codes(_eskd_drw(tt=["1. Сверлить отверстие по месту при сборке"]))
    assert "Q09" not in _eskd_codes(_eskd_drw(tt=["1. По п. 3 ТТ чертежа", "2. ГОСТ 2789-73 раздел 2"]))
    assert "Q15" in _eskd_codes(_eskd_drw(tt=["1. Острые кромки притупить."]))
    assert "Q15" not in _eskd_codes(_eskd_drw(tt=["1. Острые кромки притупить R0,3."]))
    assert "Q20" not in _eskd_codes(_eskd_drw(tt=["1. Размер 10 +0,1 -0,05", "2. Диапазон 10-12"]))


def test_eskd_brackets_and_title():
    codes = _eskd_codes(_eskd_drw(lines=[("(25)", 50, 50, 3.5), ("(⌀40)", 80, 50, 3.5), ("[⌀30H7]", 110, 50, 3.5),
                                        ("(1:20)", 140, 50, 3.5)]))
    assert "Q13" in codes and "Q14" in codes
    codes = _eskd_codes(_eskd_drw(tt=["1. Размеры в скобках — после сборки.", "2. Обработку по размерам в квадратных "
                                    "скобках производить совместно с дет. 5"],
                                lines=[("(⌀40)", 80, 50, 3.5), ("[⌀30H7]", 110, 50, 3.5)]))
    assert "Q13" not in codes and "Q14" not in codes
    assert "Q16" in _eskd_codes(_eskd_drw(mass_raw="0,25 кг"))
    assert "Q16" not in _eskd_codes(_eskd_drw(mass_raw="0,25"))
    assert "Q17" in _eskd_codes(_eskd_drw(material="Сталь 45 или Сталь 40Х"))
    assert "Q17" not in _eskd_codes(_eskd_drw(material="Круг 20 ГОСТ 2590-2006 / Сталь 45 ГОСТ 1050-2013"))
    d = _eskd_drw(designation="ДТМЛ.123456.001ГЧ", tt=["1. * Размеры для справок"])
    assert "Q19" in _eskd_codes(d)
    assert "Q19" not in _eskd_codes(_eskd_drw(designation="ДТМЛ.123456.001СБ", tt=["1. * Размеры для справок"]))


def test_eskd_views_dims_gdt_roughness():
    codes = _eskd_codes(_eskd_drw(lines=[("Вид спереди", 50, 50, 5), ("А-Б", 80, 50, 5), ("а-а", 110, 50, 5)]))
    assert "Q23" in codes and "Q24" in codes
    assert "Q24" not in _eskd_codes(_eskd_drw(lines=[("А-А", 50, 50, 5), ("Б-Б", 80, 50, 5)]))
    codes = _eskd_codes(_eskd_drw(lines=[("A-A", 50, 50, 5), ("A", 70, 80, 5)]))
    assert "Q25" in codes
    # размеры
    from checking.model import Dimension
    d = _eskd_drw(lines=[("+/-0,1", 20, 20, 3.5), ("30 град", 50, 20, 3.5), ("1x45°", 80, 20, 3.5)])
    d.dims = [Dimension("⌀20", 0, (10, 10, 20, 14), upper=0.021, lower=0.0, upper_raw="0,021", lower_raw="+0,00", symmetric=False),
              Dimension("120", 0, (10, 60, 20, 64), nominal=120.0), Dimension("120", 0, (100, 90, 110, 94), nominal=120.0)]
    codes = _eskd_codes(d)
    assert {"Q26", "Q27", "Q28", "Q29", "Q30", "Q31"} <= set(codes), codes
    d2 = _eskd_drw(); d2.dims = [Dimension("⌀20", 0, (10, 10, 20, 14), upper=0.021, lower=0.0, upper_raw="+0,021", lower_raw="0")]
    codes2 = set(_eskd_codes(d2))
    assert "Q27" not in codes2 and "Q26" in codes2            # «+0,021 / 0»: нулевое отклонение не указывают (п. 6.6)
    # рамки допусков
    codes = _eskd_codes(_eskd_drw(lines=[("⏥ 0,05 А", 20, 20, 3.5), ("⌀0 А", 20, 40, 3.5), ("◎ 0,02 А А", 20, 60, 3.5)]))
    assert "Q32" in codes and "Q34" in codes
    assert not {"Q32", "Q33", "Q34"} & set(_eskd_codes(_eskd_drw(lines=[("⏥ 0,05", 20, 20, 3.5), ("◎ ⌀0 Ⓜ А", 20, 40, 3.5)])))
    # шероховатость
    codes = _eskd_codes(_eskd_drw(tt=["1. Ra3,2", "2. Rz 20.5 мкм", "3. Неуказанная шероховатость поверхностей Ra 6,3"]))
    assert {"Q35", "Q36", "Q37", "Q39"} <= set(codes), codes
    d = _eskd_drw(); d.roughness = [{"param": "Ra", "value": 6.3, "general": True, "page": 0, "bbox": (20, 150, 40, 156), "raw": "Ra 6,3"}]
    assert "Q38" in _eskd_codes(d)
    d = _eskd_drw(); d.roughness = [{"param": "Ra", "value": 6.3, "general": True, "page": 0, "bbox": (170, 10, 190, 16), "raw": "Ra 6,3"}]
    assert "Q38" not in _eskd_codes(d)
    d = _eskd_drw()
    d.dims = [Dimension(str(i), 0, (10, 10 + i, 20, 14 + i), nominal=float(i * 10)) for i in range(1, 5)]
    assert "Q40" in _eskd_codes(d)
    d.tt = ["1. Шероховатость поверхностей — по чертежу"]
    d.facts["_tt_items"] = [(t, 0, None) for t in d.tt]
    d.findings.clear()
    assert "Q40" not in _eskd_codes(d)


def test_eskd_positions_rotated():
    d = _eskd_drw(lines=[("5", 60, 60, 7)])
    d.facts["is_assembly"] = True
    d.facts["positions"] = [{"n": 5, "page": 0, "bbox": d.lines[0].bbox, "size": 7}]
    d.lines[0].dir = (0.0, -1.0)
    assert "Q18" in _eskd_codes(d)
    d.lines[0].dir = (1.0, 0.0)
    d.findings.clear()
    assert "Q18" not in _eskd_codes(d)


def test_spec_rules_2106():
    from checking.specdoc import check_children
    out = []
    add = lambda sev, code, msg, pg=None, bb=None: out.append((code, msg))
    rows = [
        {"pos": "1", "designation": "ДТМЛ.123456.001ТУ", "name": "Сборочный чертёж", "_section": "Документация"},
        {"designation": "ДТМЛ.123456.002", "name": "деталь", "qty": 1.5, "_section": "Детали"},
        {"pos": "3", "name": "ГОСТ 7798-70 Болт М6", "qty": 4, "_section": "Стандартные изделия"},
        {"pos": "4", "name": "Лента", "qty": 2, "unit": "шт", "_section": "Материалы"},
    ]
    check_children(add, "СБ", rows)
    codes = {c for c, _ in out}
    assert {"S15", "S16", "S17", "S18", "S19", "S20"} <= codes, codes
    out.clear()
    good = [{"designation": "ДТМЛ.123456.001СБ", "name": "Сборочный чертёж", "_section": "Документация"},
            {"pos": "1", "designation": "ДТМЛ.123456.002", "name": "Корпус", "qty": 1, "_section": "Детали"},
            {"pos": "2", "name": "Болт М6 ГОСТ 7798-70", "qty": 4, "_section": "Стандартные изделия"},
            {"pos": "3", "name": "Лента", "qty": 2.5, "unit": "м", "_section": "Материалы"}]
    check_children(add, "СБ", good)
    assert not {c for c, _ in out} & {"S15", "S16", "S17", "S18", "S19", "S20"}, out


def test_ost_070_015_full():
    from checking.model import Drawing
    from checking.ost import check_ost
    drw = Drawing(path="x.pdf", filename="x.pdf")
    drw.facts["is_assembly"] = True
    drw.tt = ["1. Технические требования по ОСТ 4 Г0.070.015 ред. 1-70",
              "2. Детали расконсервировать и промыть",
              "3. Винты затянуть плотно и равномерно",
              "4. Ручка должна вращаться плавно",
              "5. Прокладки плотно уложить в паз без складок",
              "6. Вмятины вокруг заклёпок допускаются до 0,5 мм",
              "7. Допускается не более 4 радиальных разрывов у развальцованных втулок",
              "8. Осветлённые участки на гетинаксе допускаются до 2 мм",
              "9. Допускаются царапины на покрытии",
              "10. Места клёпки окрасить эмалью по РД 107.9.4002-92",
              "12. Покрыть лаком AK-113",
              "11. Отверстие сверлить по месту при сборке"]
    drw.facts["_tt_items"] = [(t, 0, None) for t in drw.tt]
    check_ost(drw, {})
    codes = {f.code for f in drw.findings}
    assert {"O19", "O20", "O21", "O22", "O23", "O24", "O25", "O26", "O27", "O28", "O29", "O30", "O31", "O32"} <= codes, codes
    d2 = Drawing(path="y.pdf", filename="y.pdf")
    d2.facts["is_assembly"] = True
    d2.tt = ["Технические требования по ОСТ 4 ГО.070.015", "Вмятины вокруг заклёпок до 0,2 мм"]
    d2.facts["_tt_items"] = [(t, 0, None) for t in d2.tt]
    check_ost(d2, {})
    assert not {f.code for f in d2.findings} & {"O20", "O26", "O30"}, [(f.code, f.message) for f in d2.findings]


def test_eskd_codes_in_catalog():
    import re as _re
    from checking import catalog
    src = open(os.path.join(ROOT, "checking", "eskd.py"), encoding="utf-8").read()
    used = set(_re.findall(r'"(Q\d\d)"', src))
    in_cat = {r[0] for r in catalog.RULES if r[0].startswith("Q")}
    assert used == in_cat, sorted(used ^ in_cat)
    ix = {e["key"]: e for e in catalog.standards_index({})}
    for key, n in (("ГОСТ Р 2.109", 20), ("ГОСТ 2.201", 1), ("ГОСТ Р 2.201", 6), ("ГОСТ 2.301", 2), ("ГОСТ 2.309", 8), ("ГОСТ 2.307", 15),
                   ("ГОСТ 2.305", 6), ("ГОСТ Р 2.106", 10), ("ГОСТ 2.308", 6), ("ГОСТ 2.302", 3)):
        assert len(ix[key]["direct"]) >= n, (key, len(ix[key]["direct"]))

# ------------------------------------------------ ГОСТ 2.310, 2.311, 2.312, 2.316, 2.403, 9.306, 1139, 1643, 2789 ----
def _j_codes(d):
    from checking import eskd2
    eskd2.check_eskd2(d, {})
    return [f.code for f in d.findings]


def test_eskd2_tt_layout():
    d = _eskd_drw(tt=["1. Шлиф. поверхность А", "2. Согласовать с Нач. цеха", "3. Покрытие Ц6.хр", "4. Твёрдость 40…45 HRC"],
                  lines=[("Технические требования", 20, 200, 5), ("Вид", 20, 100, 3.5), ("А М5:1", 80, 50, 5)])
    codes = _j_codes(d)
    assert {"J01", "J09", "J02", "J06", "J08"} <= set(codes), codes
    d = _eskd_drw(tt=["1. Покрытие Ц6.хр", "2. Окончание. Неуказанные"], lines=[("Техническая характеристика", 20, 200, 5)])
    codes = _j_codes(d)
    assert "J03" in codes and "J02" not in codes and "J01" not in codes
    ok = _eskd_drw(tt=["1. Твёрдость 40…45 HRC", "2. Неуказанные предельные отклонения", "3. Покрытие Ц6.хр"])
    assert not [c for c in _j_codes(ok) if c in ("J01", "J06", "J09")]
    # ширина, поворот, лист
    from checking.model import Drawing, TextLine
    d = _eskd_drw(pages=((210, 297), (210, 297)))
    d.tt_lines = [TextLine("1. очень длинная строка", 0, (20, 100, 220, 105), (1, 0), 3.5, (20, 105)),
                  TextLine("2. повёрнутая", 0, (30, 100, 35, 150), (0, -1), 3.5, (30, 150)),
                  TextLine("3. второй лист", 1, (20, 100, 60, 105), (1, 0), 3.5, (20, 105))]
    assert {"J04", "J05", "J07"} <= set(_j_codes(d))


def test_eskd2_heat_coating_thread_weld():
    d = _eskd_drw(tt=["1. Ц6.хр", "2. Глубина слоя 0,7…0,9 мм", "3. НRС 40…45", "4. Оцинковать", "5. Поверхность Б — 40…45 HRC",
                      "6. Покрытие Zn6.хр", "7. HRC 40…45"],
                  lines=[("⌀М20", 30, 30, 3.5), ("Двухзаходная левая", 60, 60, 3.5)])
    codes = _j_codes(d)
    assert {"J18", "J19", "J20", "J21", "J22", "J23", "J24", "J27", "J28"} <= set(codes), codes
    ok = _eskd_drw(tt=["1. Покрытие Ц6.хр", "2. h 0,7…0,9; 40…45 HRC", "3. Покрытие поверхности А — Ц6", "4. Резьба двухзаходная левая"],
                   lines=[("А", 30, 30, 5), ("М20×1,5", 60, 60, 3.5)])
    assert not [c for c in _j_codes(ok) if c in ("J18", "J19", "J20", "J21", "J22", "J23", "J24", "J27", "J28")]
    w = _eskd_drw()
    w.facts["welds"] = [{"raw": "ГОСТ 5264-80-Т3", "std": "5264", "type": "Т", "num": 3, "method": None, "leg_mark": False, "leg": None, "page": 0},
                        {"raw": "С17-⊿5", "std": None, "type": "С", "num": 17, "method": None, "leg_mark": True, "leg": 5.0, "page": 0}]
    assert {"J25", "J26"} <= set(_j_codes(w))
    w2 = _eskd_drw()
    w2.facts["welds"] = [{"raw": "Т3-⊿5-50/30", "std": None, "type": "Т", "num": 3, "method": None, "leg_mark": True, "leg": 5.0, "page": 0}]
    assert "J34" in _j_codes(w2)
    w2.facts["welds"][0]["raw"] = "Т3-⊿5-50/100"
    w2.findings.clear()
    assert "J34" not in _j_codes(w2)


def test_eskd2_splines_gears_roughness():
    d = _eskd_drw(tt=["1. Шлицы D-8×36×40×7"])
    codes = _j_codes(d)
    assert "J29" in codes and "J30" in codes
    ok = _eskd_drw(tt=["1. Шлицы D-8×36×40H7/f7×7F8/f8 ГОСТ 1139-80"])
    assert not {"J29", "J30"} & set(_j_codes(ok))
    d = _eskd_drw(); d.roughness = [{"param": "Rz", "value": 3.2, "page": 0, "bbox": (10, 10, 20, 14)}]
    assert "J33" in _j_codes(d)


def test_eskd2_gear_table():
    from checking.model import Drawing, TextLine
    d = _eskd_drw(name="Колесо зубчатое")
    assert "J10" in _j_codes(d)                                        # таблицы параметров нет

    def row(y, label, sym="", val=""):
        d2.lines.append(TextLine(label, 0, (30, y, 80, y + 3.5), (1, 0), 3.5, (30, y + 3.5)))
        if sym:
            d2.lines.append(TextLine(sym, 0, (95, y, 100, y + 3.5), (1, 0), 3.5, (95, y + 3.5)))
        if val:
            d2.lines.append(TextLine(val, 0, (110, y, 130, y + 3.5), (1, 0), 3.5, (110, y + 3.5)))
    d2 = _eskd_drw(name="Колесо зубчатое")
    row(100, "Модуль", "m", "2")
    row(105, "Число зубьев", "z", "24,5")
    row(110, "Угол наклона линии зуба", "β", "15°")
    row(115, "Исходный контур", "", "Стандартный")
    row(120, "Степень точности", "", "8")
    codes = _j_codes(d2)
    assert {"J11", "J12", "J13", "J15", "J16", "J31"} <= set(codes), codes
    d3 = _eskd_drw(name="Колесо зубчатое")
    d2 = d3
    row(100, "Модуль", "m", "2")
    row(105, "Число зубьев", "z", "24")
    row(110, "Угол наклона линии зуба", "β", "15°")
    row(114, "Направление линии зуба", "", "Правое")
    row(118, "Исходный контур", "", "ГОСТ 13755-2015")
    row(122, "Степень точности", "", "8-С ГОСТ 1643-81")
    row(126, "Длина общей нормали", "W", "30,5")
    d3.roughness = [{"param": "Ra", "value": 1.6, "page": 0, "bbox": (10, 10, 20, 14)}]
    codes = _j_codes(d3)
    assert not {"J11", "J12", "J13", "J15", "J16", "J31", "J32"} & set(codes), codes


def test_eskd2_codes_in_catalog():
    import re as _re
    from checking import catalog
    src = open(os.path.join(ROOT, "checking", "eskd2.py"), encoding="utf-8").read()
    used = set(_re.findall(r'"(J\d\d)"', src))
    in_cat = {r[0] for r in catalog.RULES if r[0].startswith("J")}
    assert used == in_cat, sorted(used ^ in_cat)
    ix = {e["key"]: e for e in catalog.standards_index({})}
    for key, n in (("ГОСТ Р 2.316", 12), ("ГОСТ 2.403", 8), ("ГОСТ 2.310", 8), ("ГОСТ 2.311", 2), ("ГОСТ 2.312", 4),
                   ("ГОСТ 5264", 2), ("ГОСТ 9.306", 3), ("ГОСТ 1139", 3), ("ГОСТ 1643", 3), ("ГОСТ 2789", 2)):
        assert len(ix[key]["direct"]) >= n, (key, len(ix[key]["direct"]))



# ------------------------------------------------ сварные швы: тип соединения, способ, толщина (Y01–Y08) ----
def _y_codes(d):
    from checking import welds, welds2
    welds.check_welds(d, {})
    welds2.check_welds2(d, {})
    return [f.code for f in d.findings]


def test_welds2_type_words_method_thickness():
    d = _eskd_drw(tt=["1. Сварные швы по ГОСТ 5264-80.", "2. Швы стыковые У4."],
                  lines=[("У4-⊿5", 50, 50, 3.5), ("тавровый", 60, 55, 3.5)])
    codes = _y_codes(d)
    assert "Y01" in codes and "Y02" in codes, codes
    ok = _eskd_drw(tt=["1. Сварные швы по ГОСТ 5264-80."], lines=[("Т3-⊿5", 50, 50, 3.5), ("тавровый", 60, 55, 3.5)])
    assert not {"Y01", "Y02", "Y03", "Y04"} & set(_y_codes(ok))
    d = _eskd_drw(tt=[], lines=[("ГОСТ 5264-80-С17-ИП", 50, 50, 3.5), ("ГОСТ 14771-76-Т3-⊿5", 50, 70, 3.5),
                                ("ГОСТ 14771-76-С2-Р", 50, 90, 3.5)])
    codes = _y_codes(d)
    assert "Y03" in codes and "Y04" in codes and "Y07" in codes, codes
    d = _eskd_drw(tt=["1. Швы по ГОСТ 5264-80"], lines=[("5264-С1", 50, 50, 3.5)])
    d.dims.append(__import__("checking.model", fromlist=["Dimension"]).Dimension(
        raw="s8", kind="thickness", nominal=8.0, page=0, bbox=(1, 1, 2, 2)))
    assert "Y05" in _y_codes(d)
    d = _eskd_drw(tt=["1. Сварные швы по ГОСТ 5264-80"], name="Рама сварная")
    assert "Y06" in _y_codes(d)


def test_welds2_codes_in_catalog():
    import re as _re
    from checking import catalog
    src = open(os.path.join(ROOT, "checking", "welds2.py"), encoding="utf-8").read()
    used = set(_re.findall(r'"(Y\d\d)"', src))
    assert used == {r[0] for r in catalog.RULES if _re.fullmatch(r"Y\d\d", r[0])}, used


# ------------------------------------------------ ГОСТ 6033, 6357, 6636, 8724, 9013, 10948, 11284, 14034 (U01–U08) ----
def test_eskd3_threads_hardness_holes():
    from checking import eskd3
    d = _eskd_drw(tt=["1. 15…18 HRC", "2. Центровые отверстия выполнить"],
                  lines=[("М12x1,25", 20, 20, 3.5), ("трубная 1/2\"", 20, 40, 3.5), ("⌀G1/2", 20, 60, 3.5),
                         ("4 отв. ⌀8,7", 20, 80, 3.5)])
    eskd3.check_eskd3(d, {})
    codes = {f.code for f in d.findings}
    assert {"U01", "U02", "U03", "U04", "U06", "U07"} <= codes, codes
    ok = _eskd_drw(tt=["1. 40…45 HRC", "2. Центровые отверстия А2,5 ГОСТ 14034-74"],
                   lines=[("М12×1,25", 20, 20, 3.5), ("G1/2-B", 20, 40, 3.5), ("4 отв. ⌀9", 20, 80, 3.5)])
    eskd3.check_eskd3(ok, {})
    assert not {f.code for f in ok.findings} & {"U01", "U02", "U03", "U04", "U06", "U07"}


def test_eskd3_codes_in_catalog():
    import re as _re
    from checking import catalog
    src = open(os.path.join(ROOT, "checking", "eskd3.py"), encoding="utf-8").read()
    assert set(_re.findall(r'"(U\d\d)"', src)) == {r[0] for r in catalog.RULES if _re.fullmatch(r"U\d\d", r[0])}


# ------------------------------------------------ пакетные правила (checking/pack.py, ext_*.py) ----
def _ext_codes(texts_, tt=(), dims=(), fn=None):
    """Коды замечаний пакетных правил для набора надписей чертежа."""
    from checking import ext_all
    from checking.dimparse import parse_dimension
    from checking.model import Dimension
    d = _eskd_drw(tt=list(tt), lines=[(t, 20, 20 + 6 * i, 3.5) for i, t in enumerate(texts_)])
    for s in dims:
        p = parse_dimension(s)
        d.dims.append(Dimension(page=0, bbox=(1, 1, 2, 2), **p))
    (fn or ext_all.check_ext)(d, {})
    return [f.code for f in d.findings]


def test_ext_threads():
    c = _ext_codes(["М14×1,25", "М35×1,5", "М25×1,5-6g", "М30×3", "М12×1,5 левая", "М16×Ph3P1,2", "М16×3(Р1,5)",
                    "М10-6g5g", "М10-3H", "М10-5g6h", "М10-6g6g", "М10-7g", "М10-6F", "М10-6g/6H", "М10-6H/6H",
                    "М10-6g-N", "М20×2-LH-5H", "М1,2-6H"],
                   tt=["Резьба 2 класса точности"])
    for code in ("I102", "I103", "I104", "I105", "I106", "I107", "I108", "I201", "I202", "I203", "I204", "I205",
                 "I206", "I207", "I208", "I209", "I210", "I211", "I212"):
        assert code in c, (code, sorted(set(c)))
    ok = _ext_codes(["М12×1,5-6g", "М10-6H/6g", "М20×2-5H-S-LH", "М16×Ph3P1,5-6H", "М8-6H"])
    assert not [x for x in ok if x[0] == "I" and x not in ("I101",)], ok
    c = _ext_codes(["Tr40×10 (P3)", "Tr40×7 (P7)", "Тr20×4", "Tr20×4-7g", "Tr20×4-6H", "Tr20×4-7e/7H", "Tr20×4 левая"])
    assert {"I301", "I302", "I303", "I304", "I305", "I306", "I307"} <= set(c), sorted(set(c))
    c = _ext_codes(["G1/2 коническая", "G1/2-B-LH", "G1/2-B-N", "G0,5", "G5/8"])
    assert {"I401", "I402", "I404"} <= set(c), sorted(set(c))


def test_ext_registry_complete():
    """Каждый код из реестра есть в каталоге и справке, а код в исходниках — в реестре."""
    import re as _re
    from checking import catalog, ext_all, rule_docs
    ext_dir = os.path.join(ROOT, "checking")
    used = set()
    for fn in os.listdir(ext_dir):
        if fn.startswith("ext_") and fn.endswith(".py"):
            used |= set(_re.findall(r'"([A-Z]\d{3})"', open(os.path.join(ext_dir, fn), encoding="utf-8").read()))
    assert used == set(ext_all.REG), sorted(used ^ set(ext_all.REG))
    cat = {r[0]: r for r in catalog.RULES}
    for code, r in ext_all.REG.items():
        assert cat[code][1] == r["group"] and code in rule_docs.RULE_DOCS


def test_ext_elements():
    c = _ext_codes(["Шпонка 8×8×40 ГОСТ 23360-78", "Шпонка 10×8×42 ГОСТ 23360-78", "Шпонка 8×7×100 ГОСТ 23360-78",
                    "Шпонка 8×7×32", "Шпонка 4-18×11×100 ГОСТ 23360-78",
                    "Отв. центр. А2,5", "Центровое отверстие А2,5", "Отв. центр. F", "Отв. центр. А2,5 шлифовать",
                    "4 отв. ⌀8,7 под винты М8", "отв. ⌀10 под М9"],
                   dims=["9N9", "8H11", "10N9", "10D10"])
    for code in ("I601", "I602", "I603", "I604", "I605", "I606", "I607", "I608", "I609", "I701", "I702", "I703",
                 "I704", "I801", "I803"):
        assert code in c, (code, sorted(set(c)))
    ok = _ext_codes(["Шпонка 2-18×11×100 ГОСТ 23360-78", "Шпонка 8×7×32 ГОСТ 23360-78",
                     "Отв. центр. А2,5 ГОСТ 14034-74", "4 отв. ⌀9 под винты М8"], dims=["10N9", "10Js9"])
    assert not [x for x in ok if x[:2] in ("I6", "I7", "I8") and x != "I609"], ok


def test_ext_series_hardness_angles():
    c = _ext_codes(["Неуказанные радиусы R2,2 max", "Неуказанные фаски 1,3×45°", "Острые кромки притупить",
                    "Неуказанные радиусы скруглений", "Конусность 1:9", "Уклон 5°", "Твёрдость 45…50",
                    "HRA 95", "HRB 120", "HRC 30…50", "HB 700", "HRC 45…50 (HB 430…480)"],
                   dims=["⌀37", "37", "⌀37H7", "35°", "22,5°", "R0,05", "R18"])
    for code in ("U101", "U102", "U103", "U104", "U105", "U201", "U203", "U206", "U205", "U301", "U302", "U303",
                 "U304", "U305", "U306", "U401", "U402", "U403", "U405", "U406", "U407"):
        assert code in c, (code, sorted(set(c)))
    ok = _ext_codes(["Неуказанные радиусы R3 max", "Фаски 1×45°", "Острые кромки притупить R0,3", "Твёрдость 45…50 HRCэ",
                     "Конусность 1:10"], dims=["⌀40", "⌀40H7", "45°", "R16"])
    assert not [x for x in ok if x[:2] in ("U1", "U2", "U3", "U4") and x not in ("U104",)], ok
    d = _eskd_drw(name="Втулка", material="Д16Т")
    d.title.material = "Д16Т"
    d.tt = ["1. HRC 40…45"]
    d.facts["_tt_items"] = [("1. HRC 40…45", 0, None)]
    from checking import ext_all
    ext_all.check_ext(d, {})
    assert "U404" in [f.code for f in d.findings]


def test_ext_gear_table():
    from checking import ext_all
    from checking.model import Drawing, TextLine
    d = _eskd_drw(name="Колесо зубчатое")

    def row(y, label, sym="", val=""):
        d.lines.append(TextLine(label, 0, (30, y, 80, y + 3.5), (1, 0), 3.5, (30, y + 3.5)))
        if sym:
            d.lines.append(TextLine(sym, 0, (95, y, 100, y + 3.5), (1, 0), 3.5, (95, y + 3.5)))
        d.lines.append(TextLine(val, 0, (110, y, 130, y + 3.5), (1, 0), 3.5, (110, y + 3.5)))
    row(100, "Модуль", "m", "2")
    row(105, "Число зубьев", "z", "14")
    row(110, "Угол наклона линии зуба", "β", "15°")
    row(115, "Делительный диаметр", "d", "29,0")
    row(120, "Диаметр вершин зубьев", "da", "40")
    row(125, "Диаметр впадин", "df", "22")
    row(130, "Высота зуба", "h", "4")
    ext_all.check_ext(d, {})
    codes = {f.code for f in d.findings}
    assert {"U501", "U502", "U503", "U504", "U505"} <= codes, codes
    d = _eskd_drw(name="Колесо")
    d.tt = ["1. Модуль m = 2,2"]
    d.facts["_tt_items"] = [("1. Модуль m = 2,2", 0, None)]
    ext_all.check_ext(d, {})
    assert "U508" in [f.code for f in d.findings]


def test_ext_general_tolerances_fits():
    c = _ext_codes(["ГОСТ 30893.1", "ГОСТ 30893.1-m и ГОСТ 30893.1-c", "±IT14/2", "ГОСТ 25670-83", "ГОСТ 30893.2-HK",
                    "ГОСТ 30893.2", "посадка с натягом H7/f7"],
                   tt=["1. Общие допуски по ГОСТ 30893.1 и ГОСТ 30893.1-m, ГОСТ 30893.1-c, ±IT14/2",
                       "2. ГОСТ 30893.2-HK и ГОСТ 30893.2-H и ГОСТ 30893.2-K", "3. ГОСТ 30893.2"],
                   dims=["⌀30H7/f7", "⌀40H7/h6", "⌀30 +0,065/0", "⌀30 +0,021/0", "⌀3500H7", "0,8h14", "⌀30H7/u7"])
    for code in ("C301", "C302", "C304", "C310", "C401", "C403", "C101", "C102", "C103", "C108", "C201", "C202",
                 "C107"):
        assert code in c, (code, sorted(set(c)))
    c = _ext_codes([], tt=["1. Неуказанные допуски по ГОСТ 30893.1-m"],
                   dims=["0,3", "40±1", "40±0,3", "2500"])
    assert {"C305", "C307", "C308"} <= set(c), sorted(set(c))
    c = _ext_codes([], tt=["1. ГОСТ 30893.1-f"], dims=["2500"])
    assert "C306" in c, c
    ok = _ext_codes([], tt=["1. Общие допуски по ГОСТ 30893.2-mK"], dims=["⌀40H7/g6"])
    assert not [x for x in ok if x[:2] in ("C1", "C3", "C4") and x != "C402"], ok


def test_ext_weld_14771():
    c = _ext_codes(["14771-С12-ИН", "14771-С6-ИП", "14771-Т3-⊿5-УП", "14771-С4-УП", "14771-Т8-ИП", "14771-У1-ИН",
                    "14771-Н1-ИН"],
                   tt=["1. Сварка в среде углекислого газа, швы по ГОСТ 14771-76", "2. Короб закрытый"])
    for code in ("Y101", "Y106", "Y109", "Y110", "Y111"):
        assert code in c, (code, sorted(set(c)))
    c = _ext_codes(["14771-Т3-ИН"], tt=["1. Сварка в среде углекислого газа по ГОСТ 14771-76"])
    assert "Y107" in c, c
    c = _ext_codes(["14771-Т3-УП"], tt=["1. Сварка по ГОСТ 14771-76"])
    assert "Y108" in c, c
    d = _eskd_drw(tt=["1. ГОСТ 14771-76. Проволока Св-08Г2С, CO2. Контроль швов визуальный"],
                  lines=[("14771-С6-ИП", 20, 20, 3.5), ("14771-У1-ИН", 20, 30, 3.5), ("14771-Т8-ИП", 20, 40, 3.5),
                         ("14771-Н1-ИН", 20, 50, 3.5)])
    d.facts["sheet_thickness"] = 10
    from checking import ext_all
    ext_all.check_ext(d, {})
    codes = {f.code for f in d.findings}
    assert {"Y102", "Y103", "Y105"} <= codes and "Y104" not in codes, codes


def test_ext_title_block():
    from checking import ext_all
    d = _eskd_drw(name="втулка.", designation="ДТМЛ.123456.001", litera="O1", sheets="1", organization="",
                  material="")
    d.title.dates = {"Разраб": "31.02.24", "Пров": "10.03.24"}
    d.title.people = {"Разраб": "Ivanov"}
    d.title.first_use = "ДТМЛ.123456.001"
    d.title.changes = [{"num": "1"}, {"num": "3"}]
    ext_all.check_ext(d, {})
    codes = {f.code for f in d.findings}
    assert {"C601", "C603", "C604", "C606", "C609", "C610", "C611", "C612", "C614"} <= codes, codes
    d = _eskd_drw(name="Корпус. Сборочный чертёж", designation="ДТМЛ.123456.001СБ", material="Сталь 45", pages=((297, 420),) * 2)
    d.title.sheets = "3"
    d.title.dates = {"Разраб": "15.03.24", "Пров": "10.03.24"}
    ext_all.check_ext(d, {})
    codes = {f.code for f in d.findings}
    assert {"C605", "C607", "C613"} <= codes and "C602" not in codes, codes     # «Сборочный чертёж» в графе 1 — норма


def test_ext_drill_and_frames():
    c = _ext_codes(["⌀9,5 под резьбу М10", "⌀8,5 под резьбу М10×1,5"])
    assert c.count("I111") == 1, c
    from checking import ext_all
    from checking.model import Drawing, TextLine
    d = _eskd_drw()
    for i, t in enumerate(["⏥ 0,05 А", "⌯ ⌀0,05 А", "⌖ ⌀0 А", "↗ 0,02 А Б В", "∥ 0,05 А А"]):
        d.lines.append(TextLine(t, 0, (20, 20 + 8 * i, 60, 25 + 8 * i), (1, 0), 3.5, (20, 25 + 8 * i)))
    ext_all.check_ext(d, {})
    codes = {f.code for f in d.findings}
    assert {"C506", "C507", "C508", "C509", "C510"} <= codes, codes


def test_ext_roughness_gears_splines_scales():
    c = _ext_codes(["Rа 3,2", "RA 6,3", "Ra 3,2 мм", "Ra 3,2  Rz 1,6", "Rz 6,3 Rmax 3,2", "Sm 0,001", "Ra 1,6, l = 1",
                    "Ra 3,2, l = 0,8", "Ra 3,2…1,6", "Ra 3,2 ±25 %", "tp 55",
                    "8-7-С ГОСТ 1643-81", "2-2-2-Н ГОСТ 1643-81", "8-7-6-Be ГОСТ 1643-81", "8-7-6-В ГОСТ 1643-81",
                    "Колесо коническое, ГОСТ 1643-81",
                    "8×36×40×7 ГОСТ 1139-80", "Д-8×36×40×7", "D-8x36x40x7", "D-8×36H12×40f7×7D9", "D-8×36×40g6/H7×7",
                    "D-8×42×48×8 лёгкой серии", "M2:1", "М1:2.5"])
    for code in ("R101", "R102", "R103", "R104", "R105", "R106", "R107", "Z101", "Z102", "Z103", "Z104",
                 "Z105", "I901", "I902", "I903", "I904", "I905", "I906", "B101", "B102"):
        assert code in c, (code, sorted(set(c)))
    ok = _ext_codes(["Ra 3,2", "Rz 12,5", "Ra 1,6…3,2", "8-7-6-Ba ГОСТ 1643-81", "D-8×36H7/g6×40×7 ГОСТ 1139-80", "М1:2,5"])
    assert not [x for x in ok if x[:2] in ("R1", "Z1", "I9", "B1")], ok
    d = _eskd_drw(format="А4", pages=((297, 210),))
    from checking import ext_all
    ext_all.check_ext(d, {})
    assert "B103" in [f.code for f in d.findings]


def test_ext_coatings_9306_and_splines_1139():
    c = _ext_codes([], tt=["1. Покрытие Цн6.хр", "2. Покрытие Ц6мкм.хр", "3. Покрытие хр.Ц6", "4. Покрытие Ц9.хр.лкп",
                           "5. Покрытие Х0,5", "6. Покрытие Ц6.хр.радужный", "7. Покрытие Ц6,хр", "8. Покрытие Ц6хр",
                           "9. По ГОСТ 9.073-77", "10. Покрытие Цинк 6 мкм хроматированное"])
    for code in ("H101", "H102", "H103", "H104", "H105", "H106", "H107", "H108", "H109", "H110"):
        assert code in c, (code, sorted(set(c)))
    ok = _ext_codes([], tt=["1. Покрытие Ц6.хр.бцв", "2. Покрытие Ц9.хр/лкп", "3. Покрытие стн.Нп612.Х",
                            "4. Покрытие Ан.Окс.тв.нхр", "5. Покрытие Хим.Фос.прм", "6. Покрытие Н9.Кд3.т.хр",
                            "7. Покрытие Ц6.фос.окс.прм", "8. Покрытие Зл1", "9. Покрытие Хим.Н3.Ср9"])
    assert not [x for x in ok if x.startswith("H1")], ok
    c = _ext_codes(["D-10×21×26×3 ГОСТ 1139-80 H7/h7 F10/h9", "ГОСТ 1139-58"], tt=["1. Шлицы накатать"])
    assert {"I907", "I910"} <= set(c), c
    c = _ext_codes(["D-8×36H7/h7×40×7F10/h9"], tt=["1. Неуказанные радиусы R1"])
    assert "I908" in c, c


def test_ext_1643_combination_rules():
    c = _ext_codes(["8-5-6-С ГОСТ 1643-81", "7-6-9-С ГОСТ 1643-81", "10-С ГОСТ 1643-81"])
    assert {"Z108", "Z109", "Z110"} <= set(c), c
    ok = _ext_codes(["8-7-6-Ba ГОСТ 1643-81", "7-C ГОСТ 1643-81", "8-9-7-A ГОСТ 1643-81", "8-N-6-Ba ГОСТ 1643-81"])
    assert not [x for x in ok if x in ("Z108", "Z109", "Z110", "Z101")], ok


def test_ext_11284_more_and_10948():
    c = _ext_codes(["4 отв. ⌀13,5 под винты", "отв. ⌀10 под заклёпки"], dims=["⌀9H12", "⌀13,5", "R3", "R0,12"])
    assert {"I804", "I805", "I806"} <= set(c), c
    ok = _ext_codes(["4 отв. ⌀9 под винты М8"], dims=["⌀9H13", "R0,12", "R63"])
    assert not [x for x in ok if x in ("I804", "I805", "I806", "D17", "D29")], ok


def test_ext_general_tol_annex_a():
    c = _ext_codes([], tt=["1. Общие допуски по ГОСТ 30893.1: H14, h14, ±IT16/2", "2. Общие допуски по ГОСТ 30893.1: H13, h13, ±IT13/2",
                           "3. Общие допуски цилиндричности — ГОСТ 30893.2-K"])
    assert {"C313", "C314", "C412"} <= set(c), c
    ok = _ext_codes([], tt=["1. Общие допуски по ГОСТ 30893.1: H14, h14, ±IT14/2. Общие допуски формы и расположения по ГОСТ 30893.2-К"],
                    dims=["7000"])
    assert not [x for x in ok if x in ("C313", "C314", "C412", "C301", "C306")], ok
    c = _ext_codes([], tt=["1. ГОСТ 30893.1-m"], dims=["7000"])
    assert "C306" not in c, c


def test_ext_spline_table_6033():
    from checking import ext_all
    from checking.model import Drawing, TextLine
    d = _eskd_drw(name="Вал шлицевой")

    def row(y, label, sym="", val=""):
        d.lines.append(TextLine(label, 0, (30, y, 80, y + 3.5), (1, 0), 3.5, (30, y + 3.5)))
        if sym:
            d.lines.append(TextLine(sym, 0, (95, y, 100, y + 3.5), (1, 0), 3.5, (95, y + 3.5)))
        d.lines.append(TextLine(val, 0, (110, y, 130, y + 3.5), (1, 0), 3.5, (110, y + 3.5)))
    row(100, "Модуль", "m", "2")
    row(105, "Число зубьев", "z", "24")
    row(110, "Номинальный диаметр", "D", "50")
    row(115, "Диаметр основной окружности", "db", "48")
    row(120, "Угол профиля", "α", "20°")
    row(125, "Шлицы эвольвентные ГОСТ 6033-80")
    ext_all.check_ext(d, {})
    codes = {f.code for f in d.findings}
    assert {"I511", "I512", "I513"} <= codes, codes


def test_ext_23360_depth_and_tool_keys():
    c = _ext_codes(["Шпонка 7×7×32 ГОСТ 23360-78"], dims=["⌀34", "10N9", "29 0/-0,1", "37,3+0,3"])
    assert "I612" in c and c.count("I613") == 2, c
    ok = _ext_codes([], dims=["⌀34", "10N9", "29 0/-0,2", "37,3+0,2"])
    assert "I613" not in ok, ok
    c = _ext_codes(["Шпонка 7×7×32 ГОСТ 23360-78 для крепления фрезы"])
    assert "I612" not in c, c


def test_ext_hardness_scales_9013():
    c = _ext_codes(["HRE 50", "HRK 120", "HRC 40…45"], tt=["1. HRG 20"])
    assert {"U408", "U409"} <= set(c), c
    ok = _ext_codes(["HRCэ 40…45", "HRE 80", "HRD 60"])
    assert not [x for x in ok if x in ("U408", "U409")], ok


def test_ext_6357_length_2311_14034():
    c = _ext_codes(["G1/2-B-15", "G/R 1/2-B", "Специальная резьба М20×1,3", "Отв. центр. А5 ГОСТ 14034-74", "Отв. центр. F М8"])
    assert {"I406", "I407", "I112", "I705"} <= set(c), c
    c = _ext_codes(["Отв. центр. F М8"], )
    d = _eskd_drw(name="Сверло", lines=[("Отв. центр. F М8", 20, 20, 3.5)])
    from checking import ext_all
    ext_all.check_ext(d, {})
    assert "I706" in [f.code for f in d.findings]
    ok = _ext_codes(["G1/2 LH-B-40", "Сп М20×1,3", "Отв. центр. А2,5 ГОСТ 14034-74"])
    assert not [x for x in ok if x in ("I406", "I407", "I112", "I705")], ok


def _mfg(name, texts_=(), tt=()):
    from checking import ext_all
    d = _eskd_drw(tt=list(tt), lines=[(t, 20, 20 + 6 * i, 3.5) for i, t in enumerate(texts_)], name=name)
    ext_all.check_ext(d, {})
    return [f.code for f in d.findings]


def test_mfg_forging_7505():
    c = _mfg("Поковка вала", tt=["1. Неуказанные штамповочные уклоны не более 15°",
                                 "2. Радиусы закруглений наружных углов R0,5"])
    for code in ("T101", "T102", "T103", "T104", "T109", "T107", "T108"):
        assert code in c, (code, c)
    ok = _mfg("Поковка вала", tt=["1. Класс точности Т4 ГОСТ 7505-89", "2. Группа стали М1", "3. Степень сложности С2",
                                  "4. Исходный индекс 12", "5. Технические требования по ГОСТ 8479",
                                  "6. Штамповочные уклоны не более 7°"])
    assert not [x for x in ok if x[0] == "T"], ok
    bad = _mfg("Поковка вала", tt=["Класс точности Т6", "Группа стали M2", "Степень сложности С5", "Исходный индекс 24",
                                   "ГОСТ 8479, ГОСТ 7505-74"])
    for code in ("T105", "T106", "T110"):
        assert code in bad, (code, bad)
    detail = _mfg("Вал", tt=["Заготовка — поковка ГОСТ 7505-89"])          # чертёж детали, а не поковки
    assert not [x for x in detail if x in ("T101", "T102", "T103", "T104", "T109")], detail


def test_mfg_casting_26645():
    assert "T201" in _mfg("Отливка корпуса", tt=["1. Литейные уклоны 3°"])
    ok = _mfg("Отливка корпуса", tt=["1. Точность отливки 8—5—4—7 См 0,8 ГОСТ 26645-85",
                                     "2. Масса 20,35—3,15—1,35—24,85 ГОСТ 26645-85"])
    assert not [x for x in ok if x[0] == "T"], ok
    bad = _mfg("Отливка корпуса", tt=["Точность отливки 8—5—7 ГОСТ 26645-85", "Точность отливки 18—12—23—7 ГОСТ 26645-85",
                                      "Точность отливки 0—5—4—7 ГОСТ 26645-85", "Точность отливки 7T—5—4—9 ГОСТ 26645-85",
                                      "Точность отливки 8—5—4—7",
                                      "Масса 20,35—3,15—1,35—24,66 ГОСТ 26645-85", "Масса 20,35 ГОСТ 26645-85"])
    for code in ("T202", "T203", "T204", "T207", "T208", "T205", "T206"):
        assert code in bad, (code, bad)
    assert "T201" in _mfg("Корпус", tt=["Отливки по ГОСТ 26645-85"])


def test_mfg_taphole_19257():
    assert "T301" in _mfg("Плита", ["⌀10 под резьбу М12"])
    assert "T301" not in _mfg("Плита", ["⌀10,2 под резьбу М12"])
    assert "T301" not in _mfg("Плита", ["⌀10 под резьбу М12×1,25"])          # мелкий шаг — другая таблица


def _spec_pair(tmp_path, with_excel=False, rows=None):
    sys.path.insert(0, os.path.join(ROOT, "tests"))
    import make_spec_pdf
    from checking.engine import check_job
    from checking.spec import Spec
    sp = make_spec_pdf.make_spec(str(tmp_path / "ДТМЛ.333333.001 - Спецификация.pdf"), rows=rows)
    asm = os.path.join(SAMPLES, "assembly", "ДТМЛ.333333.001 СБ - Кронштейн изм.1.pdf")
    excel = None
    if with_excel:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in (["Поз.", "Обозначение", "Наименование", "Кол-во", "Масса, кг"],
                  ["", "ДТМЛ.333333.001", "Кронштейн", 1, 5.78],
                  ["1", "ДТМЛ.333333.002", "Плита", 1, 1.2],
                  ["2", "ДТМЛ.333333.003", "Планка", 2, 0.5],
                  ["3", "ДТМЛ.333333.004", "Втулка", 2, 0.1],
                  ["5", "", "Болт М8-6g×20.58 ГОСТ 7798-70", 4, 0.02],
                  ["6", "", "Шайба 8 ГОСТ 11371-78", 4, 0.01]):
            ws.append(r)
        xp = str(tmp_path / "СП Кронштейн.xlsx")
        wb.save(xp)
        excel = Spec(xp)
    ds = check_file(sp, spec=excel) + check_file(asm, spec=excel)
    check_job(ds, spec=excel)
    return ds


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "assembly")), reason="нет образцов")
def test_spec_pdf_pair(tmp_path):
    ds = _spec_pair(tmp_path)
    s, a = ds
    assert s.kind == "spec" and s.facts.get("spec_document") and len(s.facts["spec_rows"]) == 5
    assert not [f for f in s.findings if f.code == "S07"]                 # документация с кодом СБ — не замечание
    a01 = [f for f in a.findings if f.code == "A01"]
    assert a01 and "Втулка" in a01[0].message and "Болт" in a01[0].message and not [f for f in a.findings if f.code == "A02"]
    assert a.facts["spec_children_src"].startswith("спецификации")
    # лишняя позиция в спецификации: чертёж без позиции 2 — A01 про позицию 2
    import make_spec_pdf
    ds2 = _spec_pair(tmp_path, rows=[r for r in make_spec_pdf.ROWS_OK if not (len(r) > 2 and r[2] == "3")])
    assert "Втулка" not in " ".join(f.message for f in ds2[1].findings if f.code == "A01")


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "assembly")), reason="нет образцов")
def test_spec_pdf_vs_excel(tmp_path):
    ds = _spec_pair(tmp_path, with_excel=True)
    s = ds[0]
    codes = {f.code: f.message for f in s.findings}
    assert "E113" in codes and "Плита" in codes["E113"]                    # наименование
    assert "E114" in codes and "ДТМЛ.333333.003" in codes["E114"]          # количество 1 ≠ 2
    assert "E115" in codes and "Болт" in codes["E115"]                     # позиция 4 ≠ 5
    assert "E112" in codes and "Шайба" in codes["E112"]                    # есть в Excel, нет в СП
    assert "E111" not in codes
    # сборочный чертёж «…001 СБ» найден в спецификации по изделию «…001» (код документа не учитывается)
    assert "E01" not in {f.code for f in ds[1].findings}
    assert ds[1].facts["spec_item"]["designation"] == "ДТМЛ.333333.001"


def test_process_documents(tmp_path):
    import openpyxl
    from checking.process import check_process_file, link_to_drawings, recognize
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in (["Маршрутная карта"], ["Деталь", "ДТМЛ.715123.901", "Вал приводной"], ["Материал", "Сталь 45 ГОСТ 1050-2013"],
              ["005", "4110", "Заготовительная", "Станок отрезной 8Б66"], ["010", "Токарная с ЧПУ", "Станок 16К20Ф3"],
              ["015", "Слесарная"], ["020", "Контрольная"]):
        ws.append(r)
    xp = str(tmp_path / "МК Вал.xlsx")
    wb.save(xp)
    d = check_process_file(xp)
    p = d.facts["process"]
    assert d.kind == "process" and p["type"] == "MK" and p["designation"] == "ДТМЛ.715123.901"
    assert [o["no"] for o in p["operations"]] == ["005", "010", "015", "020"] and "Сталь 45" in p["material"]
    assert not d.findings
    vals = check_file(os.path.join(SAMPLES, "test", "ДТМЛ.715123.901 - Вал.pdf")) if os.path.exists(
        os.path.join(SAMPLES, "test", "ДТМЛ.715123.901 - Вал.pdf")) else []
    link_to_drawings(vals + [d])
    if vals:
        assert p["drawing"] and "Вал" in p["drawing"]
    else:
        assert p["drawing"] is None
    assert recognize(["что-то", "не то"])["type"] == ""
    import docx
    doc = docx.Document()
    doc.add_paragraph("Операционная карта")
    doc.add_paragraph("ДТМЛ.715123.901")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "010", "Токарная"
    dp = str(tmp_path / "ОК.docx")
    doc.save(dp)
    assert check_process_file(dp).facts["process"]["type"] == "OK"
    junk = tmp_path / "junk.xlsx"
    wb2 = openpyxl.Workbook()
    wb2.active.append(["a", "b"])
    wb2.save(str(junk))
    assert [f.code for f in check_process_file(str(junk)).findings] == ["F08"]


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "assembly")), reason="нет образцов")
def test_api_check_assembly_with_spec_and_process(tmp_path):
    """Сборочный чертёж + спецификация КОМПАС (PDF) + маршрутная карта — одним запросом."""
    sys.path.insert(0, os.path.join(ROOT, "tests"))
    import make_spec_pdf
    import openpyxl
    import app as webapp
    sp = make_spec_pdf.make_spec(str(tmp_path / "ДТМЛ.333333.001 - Спецификация.pdf"))
    xp = str(tmp_path / "МК.xlsx")
    wb = openpyxl.Workbook()
    wb.active.append(["Маршрутная карта"])
    wb.active.append(["ДТМЛ.333333.001", "010", "Сварочная"])
    wb.save(xp)
    asm = os.path.join(SAMPLES, "assembly", "ДТМЛ.333333.001 СБ - Кронштейн изм.1.pdf")
    client = webapp.app.test_client()
    with open(asm, "rb") as fa, open(sp, "rb") as fs, open(xp, "rb") as fp:
        r = client.post("/api/v1/check", data={"file": (fa, "СБ.pdf"), "spec": (fs, "СП.pdf"), "process": (fp, "МК.xlsx"),
                                              "ocr": "0"}, content_type="multipart/form-data")
    assert r.status_code == 200, r.json
    kinds = [d["kind"] for d in r.json["drawings"]]
    assert kinds == ["pdf", "spec", "process"], kinds
    asm_d, spec_d, proc_d = r.json["drawings"]
    assert any(f["code"] == "A01" and "Втулка" in f["message"] for f in asm_d["findings"])
    assert spec_d["facts"]["spec_document"] and proc_d["facts"]["process"]["type"] == "MK"
    assert proc_d["facts"]["process"]["drawing"] == "СБ.pdf"


@pytest.mark.skipif(not os.path.isdir(os.path.join(SAMPLES, "assembly")), reason="нет образцов")
def test_upload_skips_duplicates(tmp_path, monkeypatch):
    """Одинаковые по содержимому файлы принимаются один раз — в списке файлов, внутри ZIP, в спецификациях и технологии."""
    import io
    import time
    import zipfile
    import openpyxl
    webapp, adm = _web(tmp_path, monkeypatch)
    asm = os.path.join(SAMPLES, "assembly", "ДТМЛ.333333.001 СБ - Кронштейн изм.1.pdf")
    other = os.path.join(SAMPLES, "assembly", "Р 38.02 СБ Колонна.pdf")
    data = open(asm, "rb").read()
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zz:
        zz.writestr("в архиве.pdf", data)
        zz.writestr("другой.pdf", open(other, "rb").read())
    z.seek(0)
    xl = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active.append(["Маршрутная карта"])
    wb.save(xl)
    mk = xl.getvalue()
    r = adm.post("/api/jobs", data={
        "files": [(io.BytesIO(data), "СБ.pdf"), (io.BytesIO(data), "СБ копия.pdf"), (z, "архив.zip")],
        "process": [(io.BytesIO(mk), "МК.xlsx"), (io.BytesIO(mk), "МК копия.xlsx")], "ocr": "0"},
        content_type="multipart/form-data")
    assert r.status_code == 200, r.json
    job = webapp.JOBS[r.json["job_id"]]
    assert job["total"] == 3                                   # СБ.pdf, другой.pdf из архива, одна МК
    for _ in range(300):
        if webapp.JOBS[r.json["job_id"]]["state"] != "running":
            break
        time.sleep(0.2)
    dup = [w for w in webapp.JOBS[r.json["job_id"]]["warnings"] if "дубликат" in w]
    assert len(dup) == 3 and any("СБ копия.pdf" in w for w in dup) and any("в архиве.pdf" in w for w in dup)
    assert any("МК копия.xlsx" in w for w in dup)


def test_native_cad_formats_disabled(tmp_path):
    """Родные файлы КОМПАС и Solid Edge пока не принимаются: только PDF и изображения."""
    from checking import engine
    assert not engine.CAD_FORMATS and not ({".cdw", ".spw", ".frw", ".dft"} & engine.SUPPORTED)
    for ext in (".cdw", ".spw", ".dft"):
        f = tmp_path / f"ДТМЛ.1{ext}"
        f.write_bytes(b"x")
        assert engine.collect_files([str(tmp_path)]) == []
        d = check_file(str(f))[0]
        assert [x.code for x in d.findings] == ["F05"] and "PDF" in d.findings[0].message
        f.unlink()


def _tt_drawing(items):
    from checking.model import Drawing
    d = Drawing("a.pdf", "a.pdf", kind="pdf")
    d.tt = list(items)
    d.facts["_tt_items"] = [(t, 0, None) for t in d.tt]
    return d


def test_tt_language_rules():
    """P101–P103: точка в конце пункта, типовые опечатки, пробелы у знаков препинания."""
    from checking import ext_ttlang
    d = _tt_drawing(["1 *Размер для справок.", "2 Общие допуски по ГОСТ 30893.1: H16, h16, ±",
                     "3 Маркировать шрифтом 5-Пр3 наимменование чертежа и арт.", "4 Острые кромки притупить фаской 0,5х45°",
                     "5 Размеры , допуски допуски покрытия", "6 Зазоры,допуски по наименованию."])
    ext_ttlang.check_ttlang_ext(d)
    got = [(f.code, f.message) for f in d.findings]
    p101 = [m for c, m in got if c == "P101"]
    assert len(p101) == 2 and "Острые кромки" in p101[0] and "Размеры , допуски" in p101[1]       # «±» в конце — дробь отдельной строкой
    assert [m for c, m in got if c == "P102"] == ["В ТТ слово «наимменование»: возможно, опечатка — «наименование»"]
    assert len([1 for c, _ in got if c == "P103"]) == 3                                      # пробел перед «,», «,допуски», «допуски допуски»
    assert not [1 for c, m in got if c == "P102" and "наименованию" in m]


def test_c301_quality_notation_not_limited():
    """C301: запись по квалитету («H16, h16, ±IT16/2») верна при любом номере; нет ни класса, ни квалитета — замечание."""
    from checking import ext_tol
    for tt, bad in (("2 Общие допуски по ГОСТ 30893.1: H16, h16, ± IT16/2.", False), ("Общие допуски по ГОСТ 30893.1-2002: H14, h14, ±IT14/2", False),
                    ("Общие допуски по ГОСТ 30893.1-m", False), ("Общие допуски по ГОСТ 30893.1", True)):
        d = _tt_drawing([tt])
        ext_tol.check_30893_1(d)
        assert ("C301" in [f.code for f in d.findings]) is bad, tt


def test_v01_marking_letters_near_tt_sign():
    """V01: буквы у выноски маркирования «п.3» (Ч, г — ГОСТ 2.314-68) не обозначения изображений; V и X теперь запрещены."""
    from checking import views
    from checking.model import Drawing, PageInfo, TextLine
    def line(t, x, y, d=(1.0, 0.0)):
        return TextLine(t, 0, (x, y, x + 6, y + 5), d, 3.5, (x, y + 5))
    for letters, sign, expect in (([line("Ч", 85, 50, (0.43, 0.9))], True, False), ([line("Ч", 85, 50, (0.43, 0.9))], False, True),
                                  ([line("V", 100, 100)], False, True)):
        d = Drawing("a.pdf", "a.pdf", kind="pdf")
        d.pages = [PageInfo(0, 297, 210)]
        d.lines = letters + ([line("п.3", 76, 42)] if sign else [])
        views.check_views(d, {})
        assert ("V01" in [f.code for f in d.findings]) is expect, (letters[0].text, sign)


def test_j06_order_by_2023_edition():
    """J06: порядок групп по ГОСТ Р 2.316-2023 п. 6.5 — поверхности (кромки) раньше маркирования."""
    from checking import eskd2
    d = _tt_drawing(["1 Маркировать шрифтом 5-Пр3.", "2 Острые кромки притупить."])
    eskd2.check_tt_layout(d)
    j06 = [f for f in d.findings if f.code == "J06"]
    assert j06 and "ГОСТ Р 2.316-2023 п. 6.5" in j06[0].message
    d2 = _tt_drawing(["1 Острые кромки притупить.", "2 Маркировать шрифтом 5-Пр3."])
    eskd2.check_tt_layout(d2)
    assert not [f for f in d2.findings if f.code == "J06"]


def test_norms_module(tmp_path, monkeypatch):
    """Нормативы: список из поставки, ссылка «ГОСТ Р 2.316-2023 п. 6.5» ведёт на страницу, старая редакция — никуда."""
    from checking import norms
    monkeypatch.setattr(norms, "USER", str(tmp_path / "std"))
    items = {x["designation"]: x for x in norms.listing()}
    assert {"ГОСТ 2.314-68", "ГОСТ Р 2.316-2023", "ГОСТ 2.307-2011", "ГОСТ Р 2.201-2023", "ГОСТ Р 2.109-2023"} <= set(items)
    r = norms.lookup("ГОСТ Р 2.316-2023 п. 6.5")
    assert r["id"] == "ГОСТ Р 2.316-2023.pdf" and r["page"] and "6.5" in norms.page_text(r["id"], r["page"])
    assert norms.lookup("ГОСТ 2.316-2008 п. 4.5") is None                      # другая редакция: файла нет
    assert norms.lookup("ГОСТ 2.307-2011 п. 4.3")["page"]
    assert any("справоч" in x["snippet"].lower() for x in norms.search("размеры отмечают знаком"))
    assert norms.delete("ГОСТ 2.314-68.pdf") is False                           # файлы поставки удалить нельзя


def test_norms_api(tmp_path, monkeypatch):
    import io
    from checking import norms
    monkeypatch.setattr(norms, "USER", str(tmp_path / "std"))
    webapp, adm = _web(tmp_path, monkeypatch)
    pdf = open(os.path.join(ROOT, "standards", "ГОСТ 2.314-68.pdf"), "rb").read()
    r = adm.post("/api/admin/norms", data={"file": (io.BytesIO(pdf), "Копия 2.314.pdf"), "designation": "ГОСТ 2.314-68 копия", "title": "т"},
                 content_type="multipart/form-data")
    assert r.status_code == 200 and r.json["source"] == "user"
    assert any(x["id"] == "Копия 2.314.pdf" for x in adm.get("/api/admin/norms").json)
    assert adm.get("/norms/" + "Копия 2.314.pdf").mimetype == "application/pdf"
    assert adm.post("/api/admin/norms", data={"file": (io.BytesIO(b"text"), "x.pdf")}, content_type="multipart/form-data").status_code == 400
    assert adm.get("/api/admin/norms/lookup?std=ГОСТ Р 2.316-2023 п. 6.5").json["page"]
    assert adm.delete("/api/admin/norms/Копия 2.314.pdf").json["deleted"]
    assert adm.delete("/api/admin/norms/ГОСТ 2.314-68.pdf").status_code == 400
    adm.post("/api/admin/users", json={"name": "ivanov", "password": "pw1234"})
    u = webapp.app.test_client()
    u.post("/api/login", json={"name": "ivanov", "password": "pw1234"})
    assert u.get("/api/admin/norms").status_code == 403 and u.get("/norms/ГОСТ 2.314-68.pdf").status_code == 403


def test_data_folder_migration_and_backup(tmp_path):
    """Данные exe лежат в «data»: прежние файлы рядом с exe переносятся один раз, при обновлении exe их не затирают; есть резервные копии."""
    from checking import config
    root, data = tmp_path / "app", tmp_path / "app" / "data"
    root.mkdir()
    (root / "config.json").write_text('{"a": 1}', encoding="utf-8")
    (root / "users.json").write_text("[]", encoding="utf-8")
    assert sorted(config.migrate_legacy(str(root), str(data))) == ["config.json", "users.json"]
    (data / "config.json").write_text('{"a": 2}', encoding="utf-8")                   # правка пользователя
    (root / "config.json").write_text('{"a": 99}', encoding="utf-8")                  # старый файл вернулся при копировании сборки
    assert config.migrate_legacy(str(root), str(data)) == []
    assert (data / "config.json").read_text(encoding="utf-8") == '{"a": 2}'
    b1 = config.backup_data(str(data))
    assert b1 and (tmp_path / "app" / "data" / "backup").exists()
    assert config.backup_data(str(data)) is None                                     # ничего не менялось
    (data / "config.json").write_text('{"a": 3}', encoding="utf-8")
    import time
    time.sleep(1.1)
    assert config.backup_data(str(data), keep=1)
    assert len(os.listdir(data / "backup")) == 1


def test_spec_qty_and_sections_gost_r_2106_2019():
    """ГОСТ Р 2.106-2019: «Кол.» необязательна у прочих и программных изделий, у материалов допустимы «Техн.»/«Приб.»,
    у комплектов обязательна; раздел «Программные изделия и базы данных» — между «Деталями» и «Стандартными изделиями»."""
    from checking.specdoc import SECTIONS, canon_section, check_children
    assert canon_section("Комплекты") == "Комплекты" and canon_section("Комплексы") == "Комплексы"
    assert SECTIONS.index("Детали") < SECTIONS.index("Программные изделия и базы данных") < SECTIONS.index("Стандартные изделия")
    found: dict[str, str] = {}
    rows = [
        {"_section": "Детали", "name": "Корпус", "designation": "АБВГ.123456.001", "qty": 1.0, "pos": "1"},
        {"_section": "Программные изделия и базы данных", "name": "Прошивка", "qty": None, "pos": "2"},
        {"_section": "Стандартные изделия", "name": "Болт ГОСТ 7798-70", "qty": 4.0, "pos": "3"},
        {"_section": "Прочие изделия", "name": "Кабель", "qty": None, "pos": "4"},
        {"_section": "Материалы", "name": "Припой", "qty": None, "qty_raw": "Техн.", "unit": "", "pos": "5"},
        {"_section": "Материалы", "name": "Лак", "qty": None, "qty_raw": "", "unit": "", "pos": "6"},
        {"_section": "Комплекты", "name": "ЗИП", "qty": None, "pos": "7"},
    ]
    check_children(lambda sev, code, msg, pg, bb: found.setdefault(code, msg), "", rows)
    assert "S01" not in found
    msg = found["S09"]
    assert "Лак" in msg and "ЗИП" in msg
    assert not any(x in msg for x in ("Прошивка", "Кабель", "Припой"))


# ---------------------------------------------------------------- профиль оборудования и правила T4xx
def _eq_drawing(**kw):
    """Чертёж без файла: только данные, нужные правилам профиля оборудования."""
    from checking.model import Dimension, Drawing
    d = Drawing("x.pdf", "x.pdf")
    d.title.found = True
    d.title.material = kw.get("material", "")
    if kw.get("sheet"):
        d.facts["sheet_thickness"] = kw["sheet"]
    for n, vertical, kind in kw.get("dims", []):
        d.dims.append(Dimension(raw=str(n), page=0, bbox=(0, 0, 1, 1), kind=kind, nominal=float(n), vertical=vertical))
    d.tt = kw.get("tt", [])
    d.facts["gear"] = kw.get("gear", [])
    d.facts["limits"] = kw.get("limits", [])
    d.roughness = kw.get("roughness", [])
    for raw, up, lo in kw.get("tol", []):
        d.dims.append(Dimension(raw=raw, page=0, bbox=(0, 0, 1, 1), kind="diameter", nominal=40.0, upper=up, lower=lo))
    return d


def _eq_codes(d, items=None):
    import copy
    from checking import equipment
    from checking.ext_equip import check_equip_ext
    d = copy.deepcopy(d)                                           # один чертёж проверяем с разными профилями
    cfg = {"equipment": {"items": equipment.clean_items(items)}} if items is not None else {}
    check_equip_ext(d, cfg)
    return sorted({f.code for f in d.findings})


def test_equipment_default_profile_from_erp():
    from checking import equipment as E
    its = {i["id"]: i for i in E.default_items()}
    assert len(its) == 104 and all(i["kind"] != "other" for i in its.values())
    assert its[29]["kind"] == "lathe" and its[29]["params"]["swing"] == 400          # 16К20: высота центров 200 мм
    assert its[3]["params"]["swing"] == 1000 and its[4]["params"]["swing"] == 630    # 1М65, 1М63Б
    assert its[51]["params"]["drill_d"] == 35 and its[76]["params"]["drill_d"] == 25  # 2Н135, 2С125-04
    assert its[47]["params"]["gear_d"] == 500 and its[18]["params"]["gear_d"] == 320
    laser = its[63]
    assert laser["kind"] == "laser_sheet" and laser["params"]["laser_kw"] == 20
    assert (laser["params"]["bed_l"], laser["params"]["bed_w"]) == (12000, 3000) and laser["src"]["bed_l"] == "manual"   # 3×12 м
    assert its[64]["params"]["bed_l"] == 6000 and its[64]["src"]["bed_l"] == "model"
    assert laser["src"]["t_steel"] == "estimate"                                      # ориентир, а не паспорт
    assert its[61]["kind"] == "rolls" and its[61]["params"]["roll_w"] == 2500
    assert its[11]["kind"] == "lathe_cnc" and its[15]["kind"] == "lathe_carousel" and its[32]["kind"] == "edm_wire"
    assert E.clean_items([{"type": "Станок токарный", "model": "X", "params": {"swing": "400,5", "foo": 1}}])[0]["params"] == {"swing": 400.5}


def test_equip_rules_sheet_and_laser():
    big = _eq_drawing(sheet=60, material="Лист Ст3сп", dims=[(13000, False, "linear"), (1800, True, "linear")])
    assert _eq_codes(big) == ["T401", "T402"]                      # поле лазеров до 12 000 мм; сталь — до 40 мм
    long_ok = _eq_drawing(sheet=8, material="Лист Ст3сп", dims=[(11500, False, "linear"), (2500, True, "linear")])
    assert _eq_codes(long_ok) == []                                # 20-киловаттный лазер берёт лист 3×12 м
    wide = _eq_drawing(sheet=8, material="Лист Ст3сп", dims=[(4200, False, "linear"), (4000, True, "linear")])
    assert _eq_codes(wide) == []                                   # ширину по размерам не оцениваем: размеры повторяются на видах
    small = _eq_drawing(sheet=8, material="Лист Ст3сп", dims=[(100, True, "linear"), (55, False, "linear")])
    assert _eq_codes(small) == []
    thin_steel = _eq_drawing(sheet=45, material="Лист 12Х18Н10Т", dims=[(300, False, "linear"), (200, True, "linear")])
    assert _eq_codes(thin_steel) == ["T402"]                       # нержавейка: у лучшего лазера предел 30 мм
    only_small = [{"type": "Станок лазерной резки металла, 1,5кW", "model": "USI-3015L"}]
    assert "T402" in _eq_codes(_eq_drawing(sheet=10, material="Лист Ст3сп", dims=[(300, False, "linear"), (200, True, "linear")]), only_small)
    no_laser = [{"type": "Станок токарно-винторезный", "model": "16К20"}]
    assert _eq_codes(big, no_laser) == []                          # нет лазеров — правило молчит


def test_equip_rules_turning_gears_processes():
    rubber = _eq_drawing(material="Резина листовая ТМКЩ-С-3 ГОСТ 7338-90", dims=[(1250, False, "diameter"), (2000, False, "linear")])
    assert _eq_codes(rubber, [{"type": "Станок токарно-винторезный", "model": "1М65"}]) == []   # резину не точат
    paspport = _eq_drawing(dims=[(1250, False, "diameter")])
    paspport.title.found = False
    assert _eq_codes(paspport, [{"type": "Станок токарно-винторезный", "model": "1М65"}]) == []  # не КД (паспорт, инструкция)
    shaft = _eq_drawing(dims=[(1250, False, "diameter"), (2000, False, "linear")])
    assert _eq_codes(shaft) == []                                  # карусельные станки без диаметра: деталь может пойти на них
    lathes = [{"type": "Станок токарно-винторезный", "model": "1М65"}]
    assert _eq_codes(shaft, lathes) == ["T403"]
    gear = [{"m": "6", "z": "100", "d": "600", "spline": False}]
    assert _eq_codes(_eq_drawing(gear=gear)) == []                 # у зубодолбёжных станков диаметр не задан
    hob = [{"type": "Станок зубофрезерный", "model": "53А50Н"}]
    assert _eq_codes(_eq_drawing(gear=gear), hob) == ["T404"]
    assert _eq_codes(_eq_drawing(gear=[{"m": "3", "z": "40", "d": None, "spline": False}]), hob) == []
    assert _eq_codes(_eq_drawing(gear=gear), lathes) == ["T404"]   # зубообрабатывающих станков нет совсем
    tt = _eq_drawing(tt=["Закалить, HRC 45…50", "Покрытие Ц6.хр", "Закалка ТВЧ"])
    assert _eq_codes(tt) == ["T408"]                               # печи и ТВЧ есть, гальваники нет — кооперация
    assert _eq_codes(tt, lathes) == ["T406", "T408"]
    assert _eq_codes(_eq_drawing(tt=["Зубошлифование после термообработки"]), lathes) == ["T407"]
    hard = _eq_drawing(tol=[("⌀40 +0,002 / -0,002", 0.002, -0.002)], roughness=[{"param": "Ra", "value": 0.1}])
    assert _eq_codes(hard) == ["T405"]
    assert _eq_codes(_eq_drawing(tol=[("⌀58 -0,05 / -0,06", -0.05, -0.06)], roughness=[{"param": "Ra", "value": 0.4}])) == []
    assert _eq_codes(_eq_drawing(limits=[{"dim": "10 х1", "grade": 1}])) == []   # «х1» из разбора — не допуск
    assert _eq_codes(hard, lathes + [{"type": "Станок хонинговальный", "model": "3К833"}]) == []


def test_equipment_api(tmp_path, monkeypatch):
    webapp, c = _web(tmp_path, monkeypatch)
    v = c.get("/api/admin/equipment").json
    assert len(v["items"]) == 104 and not v["custom"] and "laser_sheet" in v["kinds"] and "swing" in v["params"]
    items = v["items"][:3]
    items[0]["params"]["swing"] = 450
    items[0]["src"]["swing"] = "passport"
    r = c.post("/api/admin/equipment", json={"items": items})
    assert r.status_code == 200 and r.json["custom"] and len(r.json["items"]) == 3
    assert r.json["items"][0]["params"]["swing"] == 450 and r.json["items"][0]["src"]["swing"] == "passport"
    assert c.post("/api/admin/equipment", json={"items": []}).status_code == 400
    assert len(c.post("/api/admin/equipment/reset").json["items"]) == 104
    u = webapp.app.test_client()
    assert u.get("/api/admin/equipment").status_code == 401        # без входа закрыто
    # правила блока видны в панели, но блок по умолчанию выключен
    st = c.get("/api/admin/state").json
    assert any(g["key"] == "capab" and len(g["rules"]) == 8 for g in st["groups"])
    assert next(s for s in st["sections"] if s["key"] == "mfg")["module"]["enabled"] is False


def test_standards_registry():
    from checking import stdreg
    from checking.config import load_config
    reg = stdreg.build(load_config())
    rows = {r["designation"]: r for r in reg["rows"]}
    assert reg["summary"]["total"] == len(reg["rows"]) > 150
    assert rows["ГОСТ Р 2.106-2019"]["file"] and rows["ГОСТ Р 2.106-2019"]["file"]["same_edition"]   # PDF из standards/
    assert reg["summary"]["needed_missing"] == sum(1 for r in reg["rows"] if r["file"] is None and r["rules"] > 0)              # правила есть, PDF нет
    assert "ГОСТ 2.109-73" not in rows["ГОСТ Р 2.109-2023"]["cited_other"]                            # ссылки правил переведены на Р 2.109-2023
    assert rows["ГОСТ Р 2.109-2023"]["rules"] >= 20 and rows["ГОСТ Р 2.316-2023"]["rules"] >= 12
    assert stdreg.detect_designation("ГОСТ 25346-2013.pdf") == "ГОСТ 25346-2013"
    assert stdreg.detect_designation("scan.pdf", "ГОСТ Р 2.316-2023\nЕдиная система") == "ГОСТ Р 2.316-2023"
    assert stdreg.detect_designation("scan.pdf", "без номера") == ""


def test_split_designation():
    from checking.textnorm import split_designation
    assert split_designation("ДТМЛ.123456.001-02 СБ") == ("ДТМЛ.123456.001", "-02", "СБ")
    assert split_designation("ДТМЛ.123456.001СБ") == ("ДТМЛ.123456.001", "", "СБ")
    assert split_designation("ДТМЛ.123456.0012-01.02") == ("ДТМЛ.123456.0012", "-01.02", "")
    assert split_designation("DTML.123456.001") is None or split_designation("DTML.123456.001")[0].startswith("ДТМЛ")  # латиница приводится к кириллице
    assert split_designation("123-456789") is None


def test_strip_doc_code_any_code():
    from checking.textnorm import strip_doc_code
    assert strip_doc_code("ДТМЛ.123456.001 СБ") == "ДТМЛ.123456.001"
    assert strip_doc_code("ДТМЛ.123456.001-02ГЧ") == "ДТМЛ.123456.001-02"
    assert strip_doc_code("ДТМЛ.123456.001Э3") == "ДТМЛ.123456.001"             # код не из списка
    assert strip_doc_code("ДТМЛ.123456.001") == "ДТМЛ.123456.001"
    assert strip_doc_code("АБВГ.123.001СБ") == "АБВГ.123.001"                   # нестандартная запись: запасной путь


def test_standard_file_verification():
    from checking import stdreg
    path = os.path.join(ROOT, "standards", "ГОСТ Р 2.109-2023.pdf")
    assert stdreg.verify(path, "ГОСТ Р 2.109-2023", "ЕСКД. Основные требования к чертежам")["status"] == "ok"
    bad = stdreg.verify(path, "ГОСТ 2.307-2011", "ЕСКД. Нанесение размеров и предельных отклонений")
    assert bad["status"] == "bad" and "2.109" in bad["note"]                       # не тот ГОСТ
    year = stdreg.verify(path, "ГОСТ Р 2.109-2019")
    assert year["status"] == "bad"                                                 # не тот год
    assert stdreg.verify(path, "ГОСТ Р 2.109-2023", "Сварка. Соединения паяные основные типы")["status"] == "warn"   # не то название
    assert stdreg._key("ОСТ4 ГО.070.014") != stdreg._key("ОСТ4 Г0.070.015") and stdreg._key("ОСТ4 Г0.070.014") == stdreg._key("ОСТ4 ГО.070.014")
    assert stdreg.reconcile(path, "ГОСТ Р 2.109-2019")[0] == "ГОСТ Р 2.109-2023"       # в имени файла другой год: берём из текста
    assert stdreg._fmt(("ГОСТ", 0, "2.709", "89")) == "ГОСТ 2.709-89" and stdreg._fmt(("ГОСТ", 1, "2.109", "23")) == "ГОСТ Р 2.109-2023"


def test_standard_clauses_and_marks(tmp_path, monkeypatch):
    from checking import stdclauses
    from checking.config import load_config
    monkeypatch.setattr(stdclauses, "PATH", str(tmp_path / "marks.json"))
    d = stdclauses.detail(load_config(), "ГОСТ Р 2.109-2023.pdf")
    cl = {c["n"]: c for c in d["clauses"]}
    assert len(cl) > 100 and "5.14" in cl and cl["5.14"]["checks"] == ["Q15"]            # номер пункта на отдельной строке
    assert {"B01", "Q16"} <= set(cl["5.17"]["checks"]) and "A13" not in cl["5.17"]["checks"]
    assert "Q09" in cl["4"]["child_checks"] and not d["stale"]
    assert stdclauses._nums(" п. 9.1, 9.9") == ["9.1", "9.9"] and stdclauses._nums(" приложение А") == []
    m = stdclauses.add_mark("admin", "ГОСТ Р 2.109-2023.pdf", "ГОСТ Р 2.109-2023", "6.8", "Детали, у которых…", "проверять")
    assert stdclauses.detail(load_config(), "ГОСТ Р 2.109-2023.pdf")["clauses"][[c["n"] for c in d["clauses"]].index("6.8")]["marks"][0]["id"] == m["id"]
    with pytest.raises(ValueError):
        stdclauses.add_mark("admin", "ГОСТ Р 2.109-2023.pdf", "ГОСТ Р 2.109-2023", "6.8", "", "")             # дубль открытой заявки
    assert stdclauses.update_mark(m["id"], "done", None, "Q99")["status"] == "done"
    assert stdclauses.delete_mark(m["id"]) and stdclauses.load_marks() == []


def test_title_block_graphs_and_link_levels():
    from checking import stdclauses, std_links
    from checking.config import load_config
    d = stdclauses.detail(load_config(), "ГОСТ 2.104-2023 - Основные надписи.pdf")
    gr = {c["n"]: c for c in d["clauses"] if c["n"].startswith("гр.")}
    assert len(gr) == 41 and {"C604", "C605"} <= set(gr["гр. 8"]["checks"]) and gr["гр. 4"]["checks"] == ["C603"]
    old = stdclauses.detail(load_config(), "ГОСТ 2.104-2006 - Основные надписи.pdf")
    assert not [c for c in old["clauses"] if c["n"].startswith("гр.")]               # таблица другого вида не разбирается
    assert stdclauses._graf_nums(" гр. 14–17") == ["гр. 14", "гр. 15", "гр. 16", "гр. 17"] and stdclauses._graf_nums(" гр. 7, 8") == ["гр. 7", "гр. 8"]
    assert std_links.level("A13")[0] == "indirect" and std_links.level("Q16")[0] == "exact" and std_links.level("L01")[0] == "unreviewed"
    assert d["codes"]["C613"]["link"] == "indirect"


def test_litera_and_sheets_rules():
    from checking.ext_title import check_title_ext
    from checking.model import Drawing, PageInfo

    def codes(litera, sheets="1"):
        d = Drawing(path="x.pdf", filename="x.pdf")
        d.pages = [PageInfo(0, 210, 297)]
        d.title.found = True
        d.title.litera, d.title.sheets = litera, sheets
        check_title_ext(d)
        return [f.code for f in d.findings]

    assert "C616" in codes("Л") and "C616" in codes("О0Х")
    assert not [c for c in codes("О1") if c == "C616"] and not [c for c in codes("А Б") if c == "C616"] and not [c for c in codes("") if c == "C616"]
    assert "C617" in codes("О", sheets="") and "C617" not in codes("О", sheets="1")


def test_classifier_data_and_rules():
    from checking import classifier
    from checking.ext_class import check_classifier
    from checking.model import Drawing, PageInfo, Dimension
    assert classifier.available() and classifier.data()["version"] == "1.79.100"
    assert classifier.lookup("741211") == "Без отв." and classifier.full_name("741211").count("›") == 4
    assert classifier.features("741214")["holes"] is True and classifier.features("741211")["holes"] is False
    assert classifier.features("741514")["slots"] is True and classifier.features("715122")["central_hole"] is True
    assert classifier.children("")[0]["code"] == "04" and classifier.search("741211")[0]["code"] == "741211"

    def run(des, dims=(), lines=()):
        d = Drawing(path="x.pdf", filename="x.pdf")
        d.pages = [PageInfo(0, 210, 297)]
        d.title.found = True
        d.title.designation = des
        d.dims = list(dims)
        from checking.model import TextLine
        d.lines = [TextLine(t, 0, (10, 10, 30, 14), (1, 0), 3.5, (10, 14)) for t in lines]
        check_classifier(d)
        return [f.code for f in d.findings]

    assert run("ДТМЛ.741211.001") == ["K102"]                                   # код есть, согласуется
    assert "K101" in run("ДТМЛ.741219.001") and "K102" not in run("ДТМЛ.741219.001")   # вида 9 в группе нет
    assert run("ДТМЛ.713000.001")[0] == "K102"                                  # нули в конце — уровень не указан
    assert run("123-456789") == []                                              # другая система обозначений
    assert "K103" in run("ДТМЛ.741211.001", lines=["4 отв. ⌀10"])                # код «без отв.», на чертеже отверстия
    assert "K103" in run("ДТМЛ.741214.001")                                     # код «с отв.», отверстий нет
    assert "K105" not in run("ДТМЛ.741211.001", lines=["в диапазоне 25..35"])    # «диапазон» не паз


def test_classifier_module_off_by_default(tmp_path):
    from checking import catalog
    assert catalog.module_enabled({}, "cls") is False
    assert catalog.rule_settings({}, "K101")[0] is False                        # выключен блок — выключены все проверки
    assert catalog.rule_settings({"modules": {"cls": True}}, "K101")[0] is True
    assert catalog.GROUP_MODULE["classifier"] == "cls"


def test_doctype_classify(tmp_path):
    import openpyxl
    import pymupdf
    from checking import doctype
    x = tmp_path / "mk.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["МАРШРУТНАЯ КАРТА", "ЮМЛТ.741211.022"])
    wb.save(x)
    assert doctype.classify(str(x))[0] == doctype.PROCESS
    s = tmp_path / "sp.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Обозначение", "Наименование", "Количество"])
    wb.save(s)
    assert doctype.classify(str(s))[0] == doctype.SPEC

    def pdf(name, text):
        p = tmp_path / name
        d = pymupdf.open()
        d.new_page().insert_text((50, 50), text, fontname="helv")
        d.save(p)
        return str(p)
    # шрифт helv без кириллицы: проверяем через латиницу в регулярном выражении не получится, поэтому берём реальный PDF с текстом
    assert doctype.classify(pdf("c.pdf", "scale mass"))[0] == doctype.DRAWING
    assert doctype.classify(str(tmp_path / "a.dxf"))[0] == doctype.DRAWING
    assert doctype.classify(str(tmp_path / "t.docx"))[0] == doctype.PROCESS


def test_classifier_comparison_table():
    from checking.ext_class import check_classifier
    from checking.model import Drawing, PageInfo, TextLine

    def run(des, material="", lines=()):
        d = Drawing(path="x.pdf", filename="x.pdf")
        d.pages = [PageInfo(0, 210, 297)]
        d.title.found = True
        d.title.designation = des
        d.title.material = material
        d.lines = [TextLine(t, 0, (10, 10, 30, 14), (1, 0), 3.5, (10, 14)) for t in lines]
        check_classifier(d)
        return d

    d = run("ДТМЛ.741211.001", lines=["4 отв. ⌀10"])
    rows = {r["attr"]: r for r in d.facts["classifier"]["rows"]}
    assert rows["Отверстия"]["status"] == "diff" and "K103" in [f.code for f in d.findings]
    assert d.facts["classifier"]["code"] == "741211" and d.facts["classifier"]["path"][-1]["name"] == "Без отв."
    assert [f.severity for f in d.findings if f.code == "K103"] == ["warning"]
    d = run("ДТМЛ.715122.001")
    assert "K106" in [f.code for f in d.findings]                                     # тело вращения без ⌀ на чертеже
    d = run("ДТМЛ.741211.001", material="Швеллер 8П ГОСТ 8240")
    assert "K107" in [f.code for f in d.findings]                                     # швеллер, а код не профильный
    d = run("ДТМЛ.746211.001", material="Лист 8 ГОСТ 19903")
    assert "K107" in [f.code for f in d.findings]                                     # код профильный, а заготовка — лист
    d = run("ДТМЛ.746211.001", material="Швеллер 8П ГОСТ 8240")
    assert "K107" not in [f.code for f in d.findings]
    d = run("ДТМЛ.741211.001", material="Лист 8 ГОСТ 19903")
    assert "K107" not in [f.code for f in d.findings]                                 # «профильные» из названия класса 74 не считается


def test_standard_reference_flag(tmp_path, monkeypatch):
    from checking import stdreg
    from checking.config import load_config
    monkeypatch.setattr(stdreg, "REF_PATH", str(tmp_path / "ref.json"))
    rows = lambda: {r["designation"]: r for r in stdreg.build(load_config())["rows"]}
    assert rows()["ГОСТ 2.001-2013"]["reference"] is True                  # по умолчанию справочный
    assert stdreg.set_reference("ГОСТ 2.001-2013", False) is False and rows()["ГОСТ 2.001-2013"]["reference"] is False
    assert stdreg.set_reference("ГОСТ 2.102-2013", True) is True and rows()["ГОСТ 2.102-2013"]["reference"] is True
    assert stdreg.set_reference("ГОСТ 2.001-2013", None) is True           # возврат к значению по умолчанию


def test_standard_check_ideas():
    from checking import stdclauses
    from checking.config import load_config
    d = stdclauses.suggest(load_config(), "ГОСТ Р 2.109-2023.pdf", 8)
    assert d and d["items"] and d["found"] <= d["clauses_without_checks"]
    covered = {c["n"] for c in stdclauses.detail(load_config(), "ГОСТ Р 2.109-2023.pdf")["clauses"] if c["checks"] or c["child_checks"]}
    assert not covered & {i["n"] for i in d["items"]}                       # пункты, уже покрытые проверками, не предлагаются
    assert all(i["hint"] and i["excerpt"] and i["score"] >= 4 for i in d["items"])
    assert all({"what", "where", "how", "error", "hardness"} <= set(i["plan"]) and "Как:" in i["draft"] for i in d["items"])
    full = stdclauses.suggest(load_config(), "ГОСТ Р 2.109-2023.pdf", include_all=True)
    assert len(full["items"]) == full["clauses_without_checks"] > d["found"] and any(not i["candidate"] for i in full["items"])
    assert stdclauses._limits("не менее 590 МПа и ±0,3 мм, Rz 40") == "590 МПа, ±0,3 мм, Rz 40"
    assert stdclauses.suggest(load_config(), "нет такого.pdf") is None
